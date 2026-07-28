#!/usr/bin/env python3
"""
Build the source DATABASE (Excel) behind Chapter 1 / opening statistics of
"Non-Infectious Pathology in Asian Elephants" — elephant poisoning & unnatural mortality.
Output: Chapter1_Elephant_Mortality_Database.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = '/home/user/Workshop-on-Conference/Chapter1_Elephant_Mortality_Database.xlsx'

DARK  = '1B4332'; MED = '2D6A4F'; TINT = 'D8F3DC'; TINT2 = 'EAF7EE'
WHITE = 'FFFFFF'; RED = '8B0000'

hdr_font   = Font(name='Calibri', size=11, bold=True, color=WHITE)
title_font = Font(name='Calibri', size=14, bold=True, color=WHITE)
sub_font   = Font(name='Calibri', size=10, italic=True, color=WHITE)
bold_font  = Font(name='Calibri', size=10, bold=True, color='1A1A1A')
norm_font  = Font(name='Calibri', size=10, color='1A1A1A')
note_font  = Font(name='Calibri', size=9, italic=True, color='555555')

hdr_fill   = PatternFill('solid', fgColor=MED)
dark_fill  = PatternFill('solid', fgColor=DARK)
tint_fill  = PatternFill('solid', fgColor=TINT)
tint2_fill = PatternFill('solid', fgColor=TINT2)
white_fill = PatternFill('solid', fgColor=WHITE)

thin = Side(style='thin', color='BBBBBB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left   = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = center; cell.border = border


def write_table(ws, start_row, headers, rows, widths, zebra=True):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    style_header_row(ws, start_row, len(headers))
    r = start_row + 1
    for i, row in enumerate(rows):
        fill = tint2_fill if (zebra and i % 2) else white_fill
        for j, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=j, value=val)
            cell.font = norm_font; cell.fill = fill
            cell.alignment = left if j == 1 else center
            cell.border = border
        r += 1
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return r


def title_block(ws, title, subtitle, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = title_font; c.fill = dark_fill; c.alignment = center
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.font = sub_font; c2.fill = dark_fill; c2.alignment = center
    ws.row_dimensions[2].height = 20


wb = openpyxl.Workbook()

# ══════════════════════════════════════════════════════════════════
# SHEET 1 — README
# ══════════════════════════════════════════════════════════════════
ws = wb.active; ws.title = 'README'
title_block(ws, 'CHAPTER 1 — SOURCE DATABASE',
            'Elephant Poisoning & Unnatural Mortality in India  ·  Non-Infectious Pathology in Asian Elephants', 2)
readme = [
    ('Document', 'Non-Infectious Pathology in Asian Elephants (Prof. Dr. A B Shrivastava, NDVSU Jabalpur)'),
    ('Training Program', 'Essentials for Mortality Investigation of Asian Elephant'),
    ('Organiser', 'Chhattisgarh Forest Department · 5–6 June 2026, Raigarh'),
    ('Technical Support', 'Wildlife Institute of India (WII), Dehradun'),
    ('Lab Partners', 'ICAR-IVRI & NDVSU Jabalpur'),
    ('', ''),
    ('PURPOSE', 'This workbook holds the underlying datasets used to build the opening chapter '
                '(poisoning & unnatural-mortality statistics) of the notes.'),
    ('', ''),
    ('SHEET', 'CONTENTS'),
    ('1. README', 'This sheet — description, purpose, and source notes'),
    ('2. Poisoning_StateWise', 'Elephant deaths by poisoning, state-wise, 1997–2005 (WTI Elephant Mortality Database)'),
    ('3. Unnatural_Deaths_5yr', 'National elephant deaths by unnatural cause, 5 years (Lok Sabha, July 2024)'),
    ('4. Electrocution_StateWise', 'Electrocution deaths by state, 5-year period (Lok Sabha, 2024)'),
    ('5. Toxicity_Categories', 'The three categories of elephant toxicity with examples'),
    ('6. Population_Context', 'All-India elephant population census figures'),
    ('7. Sources', 'Full source / citation list'),
    ('', ''),
    ('IMPORTANT NOTE', 'Government tallies are periodically revised. Figures reflect the cited reporting date/window. '
                       'The 1997–2005 WTI dataset counts DOCUMENTED cases only; true figures are higher (under-reporting).'),
]
r = 4
for k, v in readme:
    kc = ws.cell(row=r, column=1, value=k); kc.font = bold_font; kc.alignment = left
    vc = ws.cell(row=r, column=2, value=v); vc.font = norm_font; vc.alignment = left
    if k in ('SHEET', 'PURPOSE', 'IMPORTANT NOTE'):
        kc.fill = tint_fill; vc.fill = tint_fill
    r += 1
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 85

# ══════════════════════════════════════════════════════════════════
# SHEET 2 — Poisoning state-wise 1997-2005
# ══════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Poisoning_StateWise')
title_block(ws, 'ELEPHANT DEATHS BY POISONING (STATE-WISE)',
            'Since 1997 · Source: WTI Elephant Mortality Database', 4)
poison_rows = [
    ('Assam', '1997–2005', 29, 'Highest — intense HEC in tea-estate/agricultural landscapes'),
    ('Orissa (Odisha)', '1997–2003', 3, ''),
    ('West Bengal', '1997–2004', 3, ''),
    ('Arunachal Pradesh', '1997–2005', 2, ''),
    ('Jharkhand', '2002–2003', 2, ''),
    ('Karnataka', '1997–2005', 2, ''),
    ('Kerala', '1997–2003', 2, ''),
    ('Uttar Pradesh', '1997–2004', 2, ''),
    ('Uttaranchal (Uttarakhand)', '2001–2004', 2, ''),
    ('Meghalaya', '1997–2005', 1, ''),
    ('Tamil Nadu', '1997–2005', 1, ''),
    ('TOTAL', '1997–2005', 49, 'Documented cases only'),
]
end = write_table(ws, 4, ['State', 'Years', 'Cases', 'Remarks'],
                  poison_rows, [26, 14, 10, 46])
# bold the TOTAL row
for c in range(1, 5):
    ws.cell(row=end-1, column=c).font = bold_font
    ws.cell(row=end-1, column=c).fill = tint_fill
ws.cell(row=end+1, column=1, value='Note: Assam accounts for 29 of 49 (59%) of documented poisoning deaths.').font = note_font

# ══════════════════════════════════════════════════════════════════
# SHEET 3 — Unnatural deaths 5-year (Lok Sabha 2024)
# ══════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Unnatural_Deaths_5yr')
title_block(ws, 'ELEPHANT DEATHS BY UNNATURAL CAUSE (NATIONAL)',
            '5-year period · Source: Lok Sabha reply, July 2024 (MoEFCC)', 4)
unnat_rows = [
    ('Electrocution', 392, '74.2%', 'Dominant cause'),
    ('Train accidents', 73, '13.8%', ''),
    ('Poaching', 50, '9.5%', ''),
    ('Poisoning', 13, '2.5%', 'Assam 10, Chhattisgarh 2, West Bengal 1'),
    ('TOTAL UNNATURAL', 528, '100%', 'All India'),
]
end = write_table(ws, 4, ['Cause', 'Deaths', 'Share', 'Remarks'],
                  unnat_rows, [24, 12, 12, 44])
for c in range(1, 5):
    ws.cell(row=end-1, column=c).font = bold_font
    ws.cell(row=end-1, column=c).fill = tint_fill

# ══════════════════════════════════════════════════════════════════
# SHEET 4 — Electrocution state-wise
# ══════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Electrocution_StateWise')
title_block(ws, 'ELECTROCUTION DEATHS BY STATE',
            '5-year period · Source: Lok Sabha reply, July 2024 (MoEFCC)', 3)
elec_rows = [
    ('Odisha', 71, ''),
    ('Assam', 55, ''),
    ('Karnataka', 52, ''),
    ('Tamil Nadu', 49, ''),
    ('Chhattisgarh', 32, 'Host state of this workshop'),
    ('Jharkhand', 30, ''),
    ('Kerala', 29, ''),
    ('Other states (balance)', 74, 'Remaining of 392 national total'),
    ('NATIONAL TOTAL', 392, '74% of all unnatural deaths'),
]
end = write_table(ws, 4, ['State', 'Electrocution Deaths', 'Remarks'],
                  elec_rows, [26, 20, 40])
for c in range(1, 4):
    ws.cell(row=end-1, column=c).font = bold_font
    ws.cell(row=end-1, column=c).fill = tint_fill

# ══════════════════════════════════════════════════════════════════
# SHEET 5 — Toxicity categories
# ══════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Toxicity_Categories')
title_block(ws, 'THREE CATEGORIES OF ELEPHANT TOXICITY',
            'Framework used in Chapter 1', 3)
cat_rows = [
    ('I. Agricultural / Plant Mycotoxins', 'Fungal toxins contaminating crops that elephants raid',
     'Cyclopiazonic acid (CPA) on Kodo millet; aflatoxins'),
    ('II. Human-Wildlife Conflict Chemicals', 'Deliberately or accidentally deployed chemicals',
     'Organophosphates, carbamates (carbofuran), cyanide, urea'),
    ('III. Environmental Toxins', 'Naturally occurring toxins in the environment',
     'Cyanobacteria (blue-green algae); toxic plants (oleander); heavy metals (Pb, As)'),
]
write_table(ws, 4, ['Category', 'Description', 'Examples'], cat_rows, [34, 40, 46])

# ══════════════════════════════════════════════════════════════════
# SHEET 6 — Population context
# ══════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Population_Context')
title_block(ws, 'ALL-INDIA ELEPHANT POPULATION',
            'Census context for mortality figures', 3)
pop_rows = [
    ('2017', 29964, 'Synchronised all-India census; ~60% of global wild Asian elephants'),
    ('Chhattisgarh (approx.)', 450, 'State elephant population (workshop host state); 450+'),
]
write_table(ws, 4, ['Census / Region', 'Elephants', 'Note'], pop_rows, [26, 14, 52])

# ══════════════════════════════════════════════════════════════════
# SHEET 7 — Sources
# ══════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Sources')
title_block(ws, 'SOURCES & CITATIONS', 'Data provenance for Chapter 1', 2)
src_rows = [
    ('Poisoning state-wise 1997–2005', 'WTI (Wildlife Trust of India) Elephant Mortality Database'),
    ('Unnatural deaths (5-yr, 528)', 'Lok Sabha Unstarred Question reply, July 2024 (MoEFCC, GoI); reported via Business Standard / The Week'),
    ('Electrocution state-wise', 'Lok Sabha reply, July 2024 (MoEFCC, Government of India)'),
    ('Long-run electrocution share (~67%)', 'Project Elephant / MoEFCC-derived datasets, 2009–2025'),
    ('Population census (29,964)', 'Synchronised All-India Elephant Census 2017, Project Elephant, MoEFCC'),
    ('Toxicity categories & mechanisms', 'Prof. Dr. A B Shrivastava presentation (06-06-2026); peer-reviewed elaboration'),
    ('Kodo millet / CPA (Bandhavgarh)', 'ICAR-IVRI toxicology report (5 Nov 2024); Mongabay India; National Herald reporting'),
]
write_table(ws, 4, ['Dataset', 'Source / Citation'], src_rows, [34, 86])

# Freeze header panes on data sheets
for name in ['Poisoning_StateWise', 'Unnatural_Deaths_5yr', 'Electrocution_StateWise',
             'Toxicity_Categories', 'Population_Context', 'Sources']:
    wb[name].freeze_panes = 'A5'

wb.save(OUT)
import os
print(f'  ✓  {OUT}  ({os.path.getsize(OUT)//1024} KB)  ·  {len(wb.sheetnames)} sheets')
print('     Sheets:', ', '.join(wb.sheetnames))
