#!/usr/bin/env python3
"""Build Excel database + analytical report (PDF & DOCX) from the two divisions' death registers."""
import os
from collections import Counter, defaultdict
from elephant_death_data import R
from build_death_analysis import agg, CH

DIR='/home/user/Workshop-on-Conference'
A=agg()
def pc(n): return round(100*n/A['total'])

# ============================================================ EXCEL
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

DARK='1B4332'; MED='2D6A4F'; TINT2='EAF7EE'; WHITE='FFFFFF'; TINT='D8F3DC'; REDL='FAE0D8'
hf=Font(name='Calibri',size=10,bold=True,color=WHITE); tf=Font(name='Calibri',size=14,bold=True,color=WHITE)
sf=Font(name='Calibri',size=9,italic=True,color=WHITE); nf=Font(name='Calibri',size=9,color='1A1A1A')
bf=Font(name='Calibri',size=9,bold=True,color='1A1A1A')
hfill=PatternFill('solid',fgColor=MED); dfill=PatternFill('solid',fgColor=DARK)
t2=PatternFill('solid',fgColor=TINT2); wf=PatternFill('solid',fgColor=WHITE); tfill=PatternFill('solid',fgColor=TINT)
thin=Side(style='thin',color='CCCCCC'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ctr=Alignment('center','center',wrap_text=True); lft=Alignment('left','center',wrap_text=True)

def hdr(ws,row,n):
    for c in range(1,n+1):
        x=ws.cell(row=row,column=c); x.font=hf; x.fill=hfill; x.alignment=ctr; x.border=bd
def titleblk(ws,t,s,n):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
    c=ws.cell(1,1,t); c.font=tf; c.fill=dfill; c.alignment=ctr; ws.row_dimensions[1].height=28
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
    c=ws.cell(2,1,s); c.font=sf; c.fill=dfill; c.alignment=ctr; ws.row_dimensions[2].height=18
def table(ws,start,heads,rows,widths):
    for j,h in enumerate(heads,1): ws.cell(start,j,h)
    hdr(ws,start,len(heads)); r=start+1
    for i,row in enumerate(rows):
        fill=t2 if i%2 else wf
        for j,v in enumerate(row,1):
            x=ws.cell(r,j,v); x.fill=fill; x.border=bd; x.font=nf
            x.alignment=lft if j<=2 else ctr
        r+=1
    for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w
    return r

wb=openpyxl.Workbook()

# Sheet: Summary
ws=wb.active; ws.title='Summary'
titleblk(ws,'ELEPHANT MORTALITY ANALYSIS — RAIGARH & DHARAMJAIGARH','Combined 2020–2026 · District Raigarh, Chhattisgarh',4)
kpis=[('Total elephant deaths analysed',A['total']),
      ('Dharamjaigarh Division (2020–2026)',A['dj']),
      ('Raigarh Division (FY 2021-22 to 2026-27)',A['rg']),
      ('Unnatural deaths',f"{A['unnat']} ({pc(A['unnat'])}%)"),
      ('Natural deaths',f"{A['nat']} ({pc(A['nat'])}%)"),
      ('Electrocution (leading cause)',f"{A['elec']} ({pc(A['elec'])}%)"),
      ('Drowning / water body (2nd)',f"{A['drown']} ({pc(A['drown'])}%)"),
      ('Disease (infectious, 2026 calves)',A['disease']),
      ('Calves/juveniles (<5 yr)',f"{A['young']} ({pc(A['young'])}%)"),
      ('Adults (>15 yr)',f"{A['adult']} ({pc(A['adult'])}%)"),
      ('Legal action taken (accused booked)',f"{A['action']} (all electrocution, Dharamjaigarh)")]
r=4
for k,v in kpis:
    a1=ws.cell(r,1,k); a1.font=bf; a1.alignment=lft; a1.border=bd; a1.fill=tfill
    a2=ws.cell(r,2,v); a2.font=nf; a2.alignment=lft; a2.border=bd
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    r+=1
ws.column_dimensions['A'].width=42
for col in 'BCD': ws.column_dimensions[col].width=14

# Sheet: Line list
ws=wb.create_sheet('Line_List')
titleblk(ws,'ELEPHANT DEATH LINE-LIST (n=51)','One row per elephant · translated from divisional registers',11)
heads=['S.No.','Division','Date','Year','Range','Location','Cause Group','Cause Detail','Nat/Unnat','Sex','Age (yr)','Age Class','Legal Action','PM/Lab']
# note: 14 cols; adjust titleblk merge
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=14)
ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=14)
for j,h in enumerate(heads,1): ws.cell(4,j,h)
hdr(ws,4,len(heads)); r=5
for i,rec in enumerate(R):
    fill=t2 if i%2 else wf
    row=[i+1,rec['division'],rec['date'],rec['cyear'],rec['range'],rec['location'],
         rec['cause_group'],rec['cause_detail'],rec['nat'],rec['sex'],
         rec['age'] if rec['age'] is not None else '—',rec['age_class'],rec['action'],rec['pmlab']]
    for j,v in enumerate(row,1):
        x=ws.cell(r,j,v); x.fill=fill; x.border=bd; x.font=nf
        x.alignment=lft if j in (2,5,6,7,8,12) else ctr
        if rec['nat']=='Unnatural' and j==9: x.fill=PatternFill('solid',fgColor=REDL)
    r+=1
for j,w in enumerate([6,14,11,7,13,26,20,30,10,6,8,16,20,10],1):
    ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='A5'

# Sheet: By Year x Division
ws=wb.create_sheet('By_Year_Division')
titleblk(ws,'DEATHS BY YEAR & DIVISION','Calendar year of incident',4)
yrs=sorted({rec['cyear'] for rec in R})
rows=[]
for y in yrs:
    dj=sum(1 for rec in R if rec['cyear']==y and rec['division']=='Dharamjaigarh')
    rg=sum(1 for rec in R if rec['cyear']==y and rec['division']=='Raigarh')
    rows.append([y,dj,rg,dj+rg])
rows.append(['TOTAL',A['dj'],A['rg'],A['total']])
end=table(ws,4,['Year','Dharamjaigarh','Raigarh','Total'],rows,[12,16,12,10])
for c in range(1,5): ws.cell(end-1,c).font=bf; ws.cell(end-1,c).fill=tfill
ch=BarChart(); ch.type='col'; ch.title='Deaths by Year'; ch.style=10
data=Reference(ws,min_col=2,max_col=3,min_row=4,max_row=4+len(yrs))
cats=Reference(ws,min_col=1,min_row=5,max_row=4+len(yrs))
ch.add_data(data,titles_from_data=True); ch.set_categories(cats); ch.height=7; ch.width=14
ws.add_chart(ch,'F4')

# Sheet: By Cause
ws=wb.create_sheet('By_Cause')
titleblk(ws,'DEATHS BY CAUSE','Both divisions',4)
cc=Counter(rec['cause_group'] for rec in R)
rows=[]
for cause,n in sorted(cc.items(),key=lambda x:-x[1]):
    dj=sum(1 for rec in R if rec['cause_group']==cause and rec['division']=='Dharamjaigarh')
    rg=sum(1 for rec in R if rec['cause_group']==cause and rec['division']=='Raigarh')
    rows.append([cause,dj,rg,n])
rows.append(['TOTAL',A['dj'],A['rg'],A['total']])
end=table(ws,4,['Cause Group','Dharamjaigarh','Raigarh','Total'],rows,[30,16,12,10])
for c in range(1,5): ws.cell(end-1,c).font=bf; ws.cell(end-1,c).fill=tfill

# Sheet: Age & Sex
ws=wb.create_sheet('Age_and_Sex')
titleblk(ws,'AGE CLASS & SEX PROFILE','',4)
order=['Newborn (<1 mo)','Calf (<1 yr)','Juvenile (1–5 yr)','Sub-adult (5–15 yr)','Adult (>15 yr)']
rows=[]
for o in order:
    m=sum(1 for rec in R if rec['age_class']==o and rec['sex']=='M')
    f=sum(1 for rec in R if rec['age_class']==o and rec['sex']=='F')
    u=sum(1 for rec in R if rec['age_class']==o and rec['sex']=='U')
    rows.append([o,m,f,u,m+f+u])
tot=[sum(1 for rec in R if rec['sex']=='M'),sum(1 for rec in R if rec['sex']=='F'),sum(1 for rec in R if rec['sex']=='U')]
rows.append(['TOTAL',tot[0],tot[1],tot[2],sum(tot)])
end=table(ws,4,['Age Class','Male','Female','Unknown','Total'],rows,[20,10,10,10,10])
for c in range(1,6): ws.cell(end-1,c).font=bf; ws.cell(end-1,c).fill=tfill

# Sheet: Ranges
ws=wb.create_sheet('Range_Hotspots')
titleblk(ws,'RANGE-WISE HOTSPOTS','',3)
rc=Counter(rec['range'] for rec in R)
rows=[[k,sum(1 for rec in R if rec['range']==k and rec['cause_group']=='Electrocution'),v]
      for k,v in sorted(rc.items(),key=lambda x:-x[1])]
table(ws,4,['Range','Electrocution deaths','Total deaths'],rows,[18,20,14])

wb.save(os.path.join(DIR,'Elephant_Death_Analysis_Raigarh_Dharamjaigarh.xlsx'))
print('  ✓ Excel built')

# ============================================================ SHARED report text
findings=[
 ("Electrocution is the single largest killer.",
  f"{A['elec']} of {A['total']} deaths ({pc(A['elec'])}%) were caused by electrocution — {A['elec_dj']} in Dharamjaigarh and {A['elec_rg']} in Raigarh. It accounts for 21 of the 22 'unnatural' deaths (the 22nd being a train hit). Most Dharamjaigarh cases involved illegal live wires strung around crop fields; {A['elec_adult']} of {A['elec']} electrocution victims were prime adults (>15 yr), the reproductively most valuable animals."),
 ("Drowning in water bodies is the second-largest cause and an emerging crisis.",
  f"{A['drown']} deaths ({pc(A['drown'])}%) were drownings in dams, ponds and marshes — {A['drown_rg']} in Raigarh and {A['drown_dj']} in Dharamjaigarh. Almost all victims were calves. Department-built ponds and specific dams (Panikhet, Rabo, Ghoghra, Saraimuda, Amabada, Kera Jharia) recur in the records, indicating a man-made infrastructure hazard concentrated in the Gharghoda and Chhal ranges."),
 ("Calf and juvenile mortality dominates.",
  f"{A['young']} of {A['total']} deaths ({pc(A['young'])}%) were calves/juveniles under 5 years. Raigarh's mortality is overwhelmingly of calves (drowning, weakness, brain injury, heat stroke), whereas Dharamjaigarh loses more prime adults to electrocution — two distinct mortality signatures in adjacent divisions."),
 ("A 2026 infectious-disease cluster in Dharamjaigarh calves warrants surveillance.",
  "Within weeks in April–May 2026, three Dharamjaigarh calves (1–6 months) died of acute serohaemorrhagic pneumonia, acute bacterial hepatitis, and bacterial septicaemia respectively. Such a temporally clustered set of fatal infections in young calves should trigger active EEHV / bacterial disease surveillance, histopathology and PCR on future calf deaths."),
 ("Enforcement is recorded only for electrocution, and only in one division.",
  f"Accused were booked in {A['action']} cases — every one an electrocution in Dharamjaigarh. Raigarh's register does not capture an action column, and the single accidental 11 kV-line electrocution had no accused. This points to (a) genuine offences behind illegal-wire deaths and (b) an inconsistency in how the two divisions record enforcement."),
 ("Post-mortem coverage is complete; laboratory backing is improving.",
  "Every death carries a post-mortem reference. From 2024 onward an increasing share also carries a laboratory report ('PM & Lab'), a positive trend that this workshop's necropsy-and-sampling training directly reinforces."),
 ("Mortality peaked in 2023 (Dharamjaigarh) and 2024–25 (Raigarh).",
  "Dharamjaigarh's worst year was 2023 with 8 deaths, six of them electrocution. Raigarh's toll is concentrated in 2024–2026, driven by the calf-drowning cluster. Lightning (2 deaths, 2022) and inter-elephant fight (1) round out the natural causes."),
]
recs=[
 ("Electrocution — joint enforcement & line safety",
  ["Constitute a standing Forest–Electricity Board task force for Chhal, Dharamjaigarh, Boro, Bakaruma (DJH) and Tamnar, Gharghoda (RGH) ranges.",
   "Audit, insulate/raise or underground sagging 11 kV lines in elephant corridors; enforce statutory ground clearance.",
   "Detect and dismantle illegal hooked live-wire fences; replace with legal pulsed solar energisers.",
   "Register FIRs under WPA 1972 (Sec 51) + Electricity Act (Sec 135/161); invoke strict liability of utilities.",
   "Standardise the enforcement/action column in BOTH divisions' registers."]),
 ("Drowning — water-body engineering",
  ["Barricade or fence steep-sided department ponds/dams implicated in deaths (Panikhet, Rabo, Ghoghra, Saraimuda, Amabada, Kera Jharia).",
   "Provide gentle-slope exit ramps / graded banks so calves can climb out.",
   "Prioritise Gharghoda and Chhal ranges, where calf drownings cluster.",
   "Audit all new water structures in elephant-movement zones for calf-safety before commissioning."]),
 ("Calf health & disease surveillance",
  ["Activate EEHV / bacterial disease surveillance for calf deaths; collect fresh + fixed tissue for PCR and histopathology.",
   "Investigate the April–May 2026 Dharamjaigarh calf cluster as a potential outbreak.",
   "Train field vets in calf-specific necropsy and sampling (this workshop)."]),
 ("Data & monitoring",
  ["Merge both divisions onto one standardised register (identical fields, fixed cause categories, geo-tag, action column).",
   "Maintain the accompanying Excel database and update after every mortality.",
   "Produce an annual division-level mortality analysis for management review."]),
]
charts=[('c1_year_division.png','Fig 1 — Elephant deaths per year, by division'),
        ('c2_cause.png','Fig 2 — Deaths by cause (both divisions, n=51)'),
        ('c3_cause_by_division.png','Fig 3 — Cause profile compared across the two divisions'),
        ('c4_ageclass.png','Fig 4 — Deaths by age class'),
        ('c6_elec_age.png','Fig 5 — Electrocution victims by age class'),
        ('c5_nat_unnat.png','Fig 6 — Natural vs unnatural deaths per division'),
        ('c7_ranges.png','Fig 7 — Range-wise hotspots')]

# ============================================================ PDF
from reportlab.platypus import (SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,
                                 PageBreak,Image as RLImage,HRFlowable)
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER,TA_JUSTIFY
def build_pdf():
    W,H=A4
    doc=SimpleDocTemplate(os.path.join(DIR,'Elephant_Death_Analysis_Report.pdf'),pagesize=A4,
                          leftMargin=1.8*cm,rightMargin=1.8*cm,topMargin=1.6*cm,bottomMargin=1.6*cm)
    st=getSampleStyleSheet(); story=[]
    DK=colors.HexColor('#1B4332');MD=colors.HexColor('#2D6A4F');LT=colors.HexColor('#52B788')
    TN=colors.HexColor('#D8F3DC');T2=colors.HexColor('#EAF7EE');RD=colors.HexColor('#FAE0D8')
    WH=colors.white;MT=colors.HexColor('#555');DRD=colors.HexColor('#8B0000')
    def S(n,**k):
        s=ParagraphStyle(n,parent=st['Normal'])
        for a,v in k.items(): setattr(s,a,v)
        return s
    H1=S('H1',fontSize=12,textColor=WH,leading=16,fontName='Helvetica-Bold')
    H2=S('H2',fontSize=10.5,textColor=MD,leading=14,fontName='Helvetica-Bold',spaceBefore=6,spaceAfter=2)
    BD=S('BD',fontSize=9,leading=13,alignment=TA_JUSTIFY,spaceAfter=3)
    BL=S('BL',fontSize=9,leading=13,leftIndent=12,firstLineIndent=-9,spaceAfter=2)
    CAP=S('CAP',fontSize=7.5,textColor=MT,alignment=TA_CENTER,fontName='Helvetica-Oblique',spaceAfter=6)
    TIT=S('TIT',fontSize=17,textColor=WH,leading=22,fontName='Helvetica-Bold',alignment=TA_CENTER)
    SUB=S('SUB',fontSize=9.5,textColor=LT,leading=13,fontName='Helvetica-Oblique',alignment=TA_CENTER)
    KY=S('KY',fontSize=8.5,textColor=DK,leading=12,fontName='Helvetica-Bold')
    TD=S('TD',fontSize=8,leading=11); TH=S('TH',fontSize=8,leading=11,textColor=WH,fontName='Helvetica-Bold')
    CWID=W-3.6*cm
    def h1(t):
        story.append(Spacer(1,0.15*cm))
        tb=Table([[Paragraph(t,H1)]],colWidths=[CWID])
        tb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),DK),('TOPPADDING',(0,0),(-1,-1),6),
            ('BOTTOMPADDING',(0,0),(-1,-1),6),('LEFTPADDING',(0,0),(-1,-1),8)]))
        story.append(tb); story.append(Spacer(1,0.12*cm))
    def kyb(t,danger=False):
        tb=Table([[Paragraph(t,S('K',fontSize=8.5,textColor=(DRD if danger else DK),leading=12,fontName='Helvetica-Bold'))]],colWidths=[CWID])
        tb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),(RD if danger else TN)),('BOX',(0,0),(-1,-1),0.5,MD),
            ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LEFTPADDING',(0,0),(-1,-1),8)]))
        story.append(tb); story.append(Spacer(1,0.12*cm))
    def dtbl(heads,rows,wds):
        data=[[Paragraph(str(x),TH) for x in heads]]+[[Paragraph(str(x),TD) for x in r] for r in rows]
        t=Table(data,colWidths=[w*cm for w in wds])
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),MD),('ROWBACKGROUNDS',(0,1),(-1,-1),[WH,T2]),
            ('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#BBB')),('TOPPADDING',(0,0),(-1,-1),3),
            ('BOTTOMPADDING',(0,0),(-1,-1),3),('LEFTPADDING',(0,0),(-1,-1),4),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
        story.append(t); story.append(Spacer(1,0.2*cm))
    def img(fn,cap,w=15.5):
        p=os.path.join(CH,fn)
        if os.path.exists(p):
            im=RLImage(p,width=w*cm,height=w*cm*0.5); im.hAlign='CENTER'; story.append(im)
            story.append(Paragraph(cap,CAP))
    # TITLE
    story.append(Spacer(1,0.3*cm))
    tb=Table([[Paragraph('ANALYTICAL REPORT',SUB)],[Paragraph('Elephant Mortality Analysis',TIT)],
              [Paragraph('Raigarh &amp; Dharamjaigarh Forest Divisions · District Raigarh, Chhattisgarh',SUB)],
              [Paragraph('Based on official divisional death registers, 2020–2026 (n = 51)',SUB)]],colWidths=[CWID])
    tb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),DK),('TOPPADDING',(0,0),(-1,-1),12),
        ('BOTTOMPADDING',(0,0),(-1,-1),12),('LEFTPADDING',(0,0),(-1,-1),10)]))
    story.append(tb); story.append(Spacer(1,0.25*cm))
    kyb('Prepared for: Workshop on Essentials for Mortality Investigation of Asian Elephant · '
        'Chhattisgarh Forest Department · 5–6 June 2026, Raigarh · Technical support: WII Dehradun')
    # 1 Executive summary
    h1('1.  Executive Summary')
    story.append(Paragraph(
        f"This report analyses <b>{A['total']} elephant deaths</b> recorded by two adjacent forest divisions of "
        f"Raigarh district — <b>Dharamjaigarh ({A['dj']}, calendar years 2020–2026)</b> and "
        f"<b>Raigarh ({A['rg']}, financial years 2021-22 to 2026-27)</b>. Two causes dominate: "
        f"<b>electrocution ({A['elec']}, {pc(A['elec'])}%)</b> and <b>drowning in water bodies "
        f"({A['drown']}, {pc(A['drown'])}%)</b>, together accounting for {pc(A['elec']+A['drown'])}% of all deaths. "
        f"Of the total, <b>{A['unnat']} ({pc(A['unnat'])}%) were unnatural</b> and <b>{A['nat']} natural</b>; "
        f"<b>{A['young']} ({pc(A['young'])}%) were calves or juveniles under five years</b>. "
        "The two divisions show distinct mortality signatures — Dharamjaigarh loses prime adults to illegal-wire "
        "electrocution, while Raigarh loses calves to drowning in dams and ponds.",BD))
    dtbl(['Metric','Dharamjaigarh','Raigarh','Combined'],
         [['Total deaths',A['dj'],A['rg'],A['total']],
          ['Electrocution',A['elec_dj'],A['elec_rg'],A['elec']],
          ['Drowning / water body',A['drown_dj'],A['drown_rg'],A['drown']],
          ['Unnatural (all causes)',sum(1 for r in R if r['division']=='Dharamjaigarh' and r['nat']=='Unnatural'),
                                    sum(1 for r in R if r['division']=='Raigarh' and r['nat']=='Unnatural'),A['unnat']],
          ['Calves/juveniles (<5 yr)',sum(1 for r in R if r['division']=='Dharamjaigarh' and r['age_class'] in ('Newborn (<1 mo)','Calf (<1 yr)','Juvenile (1–5 yr)')),
                                      sum(1 for r in R if r['division']=='Raigarh' and r['age_class'] in ('Newborn (<1 mo)','Calf (<1 yr)','Juvenile (1–5 yr)')),A['young']],
          ['Legal action (accused booked)',A['action'],0,A['action']]],
         [6.2,3.6,3.4,3.6])
    # 2 Data & method
    h1('2.  Data Sources & Method')
    for t in [
        "Two official registers of the Chhattisgarh Forest Department were transcribed and translated: "
        "(a) 'Elephant Death 2020 to 2026', Dharamjaigarh Forest Division; and (b) 'Last Five Years Elephant Death', "
        "Raigarh Forest Division. Every dead elephant is one record; multi-death incidents were split into separate rows.",
        "Causes stated in the registers (per post-mortem reports) were mapped to standardised categories for comparison. "
        "Dharamjaigarh records run by calendar year; Raigarh by financial year — for temporal charts the calendar year "
        "of the incident date is used. Ages given as ranges/months were converted to a mid-point in years for age-class "
        "assignment. A small number of calf sexes were not stated in the source and are marked 'Unknown'."]:
        story.append(Paragraph(t,BD))
    story.append(PageBreak())
    # 3 Findings
    h1('3.  Key Findings')
    for i,(t,d) in enumerate(findings,1):
        story.append(Paragraph(f"<b>3.{i}  {t}</b>",H2))
        story.append(Paragraph(d,BD))
    story.append(PageBreak())
    # 4 Charts
    h1('4.  Visual Analysis')
    for fn,cap in charts:
        img(fn,cap)
    story.append(PageBreak())
    # 5 Tables
    h1('5.  Data Tables')
    story.append(Paragraph('<b>5.1  Deaths by year & division</b>',H2))
    yrs=sorted({r['cyear'] for r in R})
    dtbl(['Year','Dharamjaigarh','Raigarh','Total'],
         [[y,sum(1 for r in R if r['cyear']==y and r['division']=='Dharamjaigarh'),
              sum(1 for r in R if r['cyear']==y and r['division']=='Raigarh'),
              sum(1 for r in R if r['cyear']==y)] for y in yrs]+
         [['TOTAL',A['dj'],A['rg'],A['total']]],[4,4,4,4])
    story.append(Paragraph('<b>5.2  Deaths by cause</b>',H2))
    cc=Counter(r['cause_group'] for r in R)
    dtbl(['Cause group','Dharamjaigarh','Raigarh','Total','% '],
         [[c,sum(1 for r in R if r['cause_group']==c and r['division']=='Dharamjaigarh'),
             sum(1 for r in R if r['cause_group']==c and r['division']=='Raigarh'),n,f'{pc(n)}%']
          for c,n in sorted(cc.items(),key=lambda x:-x[1])],[6,3.2,3,2.4,2.2])
    story.append(PageBreak())
    # 6 Recommendations
    h1('6.  Recommendations')
    for i,(t,items) in enumerate(recs,1):
        story.append(Paragraph(f"<b>6.{i}  {t}</b>",H2))
        for it in items: story.append(Paragraph('•  '+it,BL))
    story.append(Spacer(1,0.2*cm))
    kyb('PRIORITY: Electrocution and calf-drowning together cause two-thirds of all deaths and are almost entirely '
        'PREVENTABLE. A Forest–Electricity Board task force on illegal wires and 11 kV sag, plus barricading/ramping '
        'of department water bodies in Gharghoda and Chhal ranges, would address the two largest killers directly.',danger=True)
    # 7 Full line list
    story.append(PageBreak()); h1('7.  Annexure — Full Line-List (n=51)')
    rows=[[i+1,r['division'][:4].upper(),r['date'],r['range'],r['cause_group'].split(' /')[0].split(' (')[0],
           r['nat'][:5],r['sex'],(r['age'] if r['age'] is not None else '—'),r['age_class'].split(' (')[0],
           ('Y' if str(r['action']).startswith('Yes') else '—')] for i,r in enumerate(R)]
    dtbl(['#','Div','Date','Range','Cause','Nat','Sex','Age','Age class','Act'],rows,
         [0.9,1.2,2.0,2.6,3.0,1.4,0.9,1.0,2.6,0.9])
    story.append(HRFlowable(width='100%',thickness=0.5,color=LT))
    story.append(Paragraph('Analytical report compiled from Dharamjaigarh &amp; Raigarh Forest Division death registers · '
        'Workshop on Mortality Investigation of Asian Elephant · Raigarh, Chhattisgarh · June 2026',CAP))
    doc.build(story)
    print('  ✓ PDF built')

# ============================================================ DOCX
from docx import Document
from docx.shared import Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
def build_docx():
    dk=RGBColor(0x1B,0x43,0x32);md=RGBColor(0x2D,0x6A,0x4F);wh=RGBColor(0xFF,0xFF,0xFF)
    bl=RGBColor(0x1A,0x1A,0x1A);mt=RGBColor(0x55,0x55,0x55);rd=RGBColor(0x8B,0,0)
    tint='D8F3DC';redl='FAE0D8';med='2D6A4F';tint2='EAF7EE'
    doc=Document(); s=doc.sections[0]; s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=Cm(1.8);s.top_margin=s.bottom_margin=Cm(1.6)
    def shade(cell,hx):
        tc=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd')
        sh.set(qn('w:val'),'clear');sh.set(qn('w:fill'),hx);tc.append(sh)
    def h1(t):
        tb=doc.add_table(rows=1,cols=1);tb.style='Table Grid';c=tb.rows[0].cells[0];shade(c,'1B4332')
        p=c.paragraphs[0];p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
        r=p.add_run(t);r.font.name='Calibri';r.font.size=Pt(13);r.font.bold=True;r.font.color.rgb=wh
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    def h2(t):
        p=doc.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(1)
        r=p.add_run(t);r.font.name='Calibri';r.font.size=Pt(10.5);r.font.bold=True;r.font.color.rgb=md
    def para(t,sz=9.5):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.space_after=Pt(3)
        r=p.add_run(t);r.font.name='Calibri';r.font.size=Pt(sz);r.font.color.rgb=bl;return p
    def bullet(t):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(0.5);p.paragraph_format.first_line_indent=Cm(-0.3)
        p.paragraph_format.space_after=Pt(1);r=p.add_run('•  '+t);r.font.name='Calibri';r.font.size=Pt(9.5);r.font.color.rgb=bl
    def keybox(t,danger=False):
        tb=doc.add_table(rows=1,cols=1);tb.style='Table Grid';c=tb.rows[0].cells[0];shade(c,redl if danger else tint)
        p=c.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
        r=p.add_run(t);r.font.name='Calibri';r.font.size=Pt(9);r.font.bold=True;r.font.color.rgb=rd if danger else dk
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    def dtable(heads,rows,widths=None):
        tb=doc.add_table(rows=1+len(rows),cols=len(heads));tb.style='Table Grid'
        for i,h in enumerate(heads):
            shade(tb.rows[0].cells[i],med);p=tb.rows[0].cells[i].paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(str(h));r.font.name='Calibri';r.font.size=Pt(8.5);r.font.bold=True;r.font.color.rgb=wh
        for ri,row in enumerate(rows):
            for ci,v in enumerate(row):
                shade(tb.rows[ri+1].cells[ci],tint2 if ri%2 else 'FFFFFF')
                p=tb.rows[ri+1].cells[ci].paragraphs[0]
                if ci not in (0,):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run(str(v));r.font.name='Calibri';r.font.size=Pt(8.5);r.font.color.rgb=bl
        if widths:
            for row in tb.rows:
                for ci,cell in enumerate(row.cells):
                    if ci<len(widths):cell.width=Cm(widths[ci])
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    def image(fn,cap,w=17):
        p=os.path.join(CH,fn)
        if os.path.exists(p):
            pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_before=Pt(4)
            pp.add_run().add_picture(p,width=Cm(w))
            cp=doc.add_paragraph();cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.space_after=Pt(6)
            r=cp.add_run(cap);r.font.name='Calibri';r.font.size=Pt(8);r.font.italic=True;r.font.color.rgb=mt
    def page_break():
        p=doc.add_paragraph();run=p.add_run();br=OxmlElement('w:br');br.set(qn('w:type'),'page');run._r.append(br)
    # TITLE
    tb=doc.add_table(rows=1,cols=1);tb.style='Table Grid';c=tb.rows[0].cells[0];shade(c,'1B4332')
    p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(16);p.paragraph_format.space_after=Pt(16)
    for txt,sz,bold,col in [('ANALYTICAL REPORT\n',9,True,RGBColor(0x52,0xB7,0x88)),
                            ('Elephant Mortality Analysis\n',18,True,wh),
                            ('Raigarh & Dharamjaigarh Forest Divisions · District Raigarh, Chhattisgarh\n',10,False,RGBColor(0xD8,0xF3,0xDC)),
                            ('Based on official divisional death registers, 2020–2026 (n = 51)',9,False,RGBColor(0xD8,0xF3,0xDC))]:
        r=p.add_run(txt);r.font.name='Calibri';r.font.size=Pt(sz);r.font.bold=bold;r.font.color.rgb=col
    doc.add_paragraph()
    keybox('Prepared for: Workshop on Essentials for Mortality Investigation of Asian Elephant · '
           'Chhattisgarh Forest Department · 5–6 June 2026, Raigarh · Technical support: WII Dehradun')
    h1('1.  Executive Summary')
    para(f"This report analyses {A['total']} elephant deaths recorded by two adjacent forest divisions of Raigarh "
         f"district — Dharamjaigarh ({A['dj']}, calendar years 2020–2026) and Raigarh ({A['rg']}, financial years "
         f"2021-22 to 2026-27). Two causes dominate: electrocution ({A['elec']}, {pc(A['elec'])}%) and drowning in "
         f"water bodies ({A['drown']}, {pc(A['drown'])}%), together {pc(A['elec']+A['drown'])}% of all deaths. "
         f"{A['unnat']} deaths ({pc(A['unnat'])}%) were unnatural and {A['nat']} natural; {A['young']} "
         f"({pc(A['young'])}%) were calves/juveniles under five years. The two divisions show distinct mortality "
         "signatures — Dharamjaigarh loses prime adults to illegal-wire electrocution, Raigarh loses calves to drowning.")
    dtable(['Metric','Dharamjaigarh','Raigarh','Combined'],
        [['Total deaths',A['dj'],A['rg'],A['total']],
         ['Electrocution',A['elec_dj'],A['elec_rg'],A['elec']],
         ['Drowning / water body',A['drown_dj'],A['drown_rg'],A['drown']],
         ['Unnatural deaths',sum(1 for r in R if r['division']=='Dharamjaigarh' and r['nat']=='Unnatural'),
          sum(1 for r in R if r['division']=='Raigarh' and r['nat']=='Unnatural'),A['unnat']],
         ['Calves/juveniles (<5 yr)',sum(1 for r in R if r['division']=='Dharamjaigarh' and r['age_class'] in ('Newborn (<1 mo)','Calf (<1 yr)','Juvenile (1–5 yr)')),
          sum(1 for r in R if r['division']=='Raigarh' and r['age_class'] in ('Newborn (<1 mo)','Calf (<1 yr)','Juvenile (1–5 yr)')),A['young']],
         ['Legal action (accused booked)',A['action'],0,A['action']]],[6,4,4,4])
    h1('2.  Data Sources & Method')
    para("Two official Chhattisgarh Forest Department registers were transcribed and translated: 'Elephant Death 2020 "
         "to 2026' (Dharamjaigarh) and 'Last Five Years Elephant Death' (Raigarh). Each dead elephant is one record; "
         "multi-death incidents were split. Register-stated causes (per post-mortem) were mapped to standardised "
         "categories. Dharamjaigarh runs by calendar year, Raigarh by financial year; temporal charts use the incident "
         "calendar year. Age ranges/months were converted to mid-point years for age-class assignment; a few calf "
         "sexes were unstated ('Unknown').")
    page_break()
    h1('3.  Key Findings')
    for i,(t,d) in enumerate(findings,1):
        h2(f'3.{i}  {t}'); para(d)
    page_break()
    h1('4.  Visual Analysis')
    for fn,cap in charts: image(fn,cap)
    page_break()
    h1('5.  Data Tables')
    h2('5.1  Deaths by year & division')
    yrs=sorted({r['cyear'] for r in R})
    dtable(['Year','Dharamjaigarh','Raigarh','Total'],
        [[y,sum(1 for r in R if r['cyear']==y and r['division']=='Dharamjaigarh'),
            sum(1 for r in R if r['cyear']==y and r['division']=='Raigarh'),
            sum(1 for r in R if r['cyear']==y)] for y in yrs]+[['TOTAL',A['dj'],A['rg'],A['total']]])
    h2('5.2  Deaths by cause')
    cc=Counter(r['cause_group'] for r in R)
    dtable(['Cause group','Dharamjaigarh','Raigarh','Total','%'],
        [[c,sum(1 for r in R if r['cause_group']==c and r['division']=='Dharamjaigarh'),
            sum(1 for r in R if r['cause_group']==c and r['division']=='Raigarh'),n,f'{pc(n)}%']
         for c,n in sorted(cc.items(),key=lambda x:-x[1])])
    page_break()
    h1('6.  Recommendations')
    for i,(t,items) in enumerate(recs,1):
        h2(f'6.{i}  {t}')
        for it in items: bullet(it)
    keybox('PRIORITY: Electrocution and calf-drowning together cause two-thirds of all deaths and are almost entirely '
           'PREVENTABLE. A Forest–Electricity Board task force on illegal wires and 11 kV sag, plus barricading/ramping '
           'of department water bodies in Gharghoda and Chhal ranges, would address the two largest killers directly.',danger=True)
    page_break()
    h1('7.  Annexure — Full Line-List (n=51)')
    dtable(['#','Div','Date','Range','Cause','Nat','Sex','Age','Age class','Act'],
        [[i+1,r['division'][:4],r['date'],r['range'],r['cause_group'].split(' /')[0].split(' (')[0],
          r['nat'][:5],r['sex'],(r['age'] if r['age'] is not None else '—'),r['age_class'].split(' (')[0],
          ('Y' if str(r['action']).startswith('Yes') else '—')] for i,r in enumerate(R)],
        [0.9,1.3,1.9,2.4,3.0,1.4,0.9,1.0,2.5,0.9])
    para('Analytical report compiled from Dharamjaigarh & Raigarh Forest Division death registers · Workshop on '
         'Mortality Investigation of Asian Elephant · Raigarh, Chhattisgarh · June 2026',sz=8)
    doc.save(os.path.join(DIR,'Elephant_Death_Analysis_Report.docx'))
    print('  ✓ DOCX built')

if __name__=='__main__':
    build_pdf(); build_docx()
    for f in ['Elephant_Death_Analysis_Report.pdf','Elephant_Death_Analysis_Report.docx',
              'Elephant_Death_Analysis_Raigarh_Dharamjaigarh.xlsx']:
        p=os.path.join(DIR,f); print(f'  {f}: {os.path.getsize(p)//1024} KB')
