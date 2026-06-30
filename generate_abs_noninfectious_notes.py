#!/usr/bin/env python3
"""
Comprehensive notes from Prof. Dr. A B Shrivastava's presentation:
"Non-Infectious Pathology in Asian Elephants"
Outputs: ABS_Elephant_NonInfectious_Notes.docx  and  .pdf
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                 TableStyle, PageBreak, Image as RLImage, HRFlowable)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY

DIR    = '/home/user/Workshop-on-Conference'
SLIDES = '/home/user/Workshop-on-Conference/abs2_slides'
OUT_D  = os.path.join(DIR, 'ABS_Elephant_NonInfectious_Notes.docx')
OUT_P  = os.path.join(DIR, 'ABS_Elephant_NonInfectious_Notes.pdf')

DARK  = RGBColor(0x1B, 0x43, 0x32); MED  = RGBColor(0x2D, 0x6A, 0x4F)
LIGHT = RGBColor(0x52, 0xB7, 0x88); TINT = RGBColor(0xD8, 0xF3, 0xDC)
TINT2 = RGBColor(0xEA, 0xF7, 0xEE); WHITE= RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A); MUTED= RGBColor(0x55, 0x55, 0x55)
REDL  = RGBColor(0xFA, 0xE0, 0xD8); RED  = RGBColor(0x8B, 0x00, 0x00)

def rgb_hex(rgb): return f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}'
def slide(n): return os.path.join(SLIDES, f'page_{n:03d}.png')

# ── DOCX HELPERS ──────────────────────────────────────────────────────────────
def shade_cell(cell, hex_color):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color); tcPr.append(shd)

def page_break(doc):
    p = doc.add_paragraph(); run = p.add_run()
    br = OxmlElement('w:br'); br.set(qn('w:type'), 'page'); run._r.append(br)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)

def sec1(doc, text):
    tbl = doc.add_table(rows=1, cols=1); tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]; shade_cell(cell, rgb_hex(DARK))
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = 'Calibri'; run.font.size = Pt(13); run.font.bold = True; run.font.color.rgb = WHITE
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def sec2(doc, text):
    tbl = doc.add_table(rows=1, cols=1); tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]; shade_cell(cell, rgb_hex(MED))
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Calibri'; run.font.size = Pt(10.5); run.font.bold = True; run.font.color.rgb = WHITE
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def sec3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.font.name = 'Calibri'; run.font.size = Pt(10); run.font.bold = True; run.font.color.rgb = MED

def para(doc, text, size=9.5, bold=False, italic=False, color=None,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.font.bold = bold
    run.font.italic = italic; run.font.color.rgb = color or BLACK
    return p

def bullet(doc, text, level=0, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(('◦  ' if level > 0 else '•  ') + text)
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.font.color.rgb = BLACK

def img(doc, n, caption='', width_cm=15):
    path = slide(n)
    if os.path.exists(path):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(6)
        cr = cp.add_run(caption); cr.font.name = 'Calibri'; cr.font.size = Pt(8)
        cr.font.italic = True; cr.font.color.rgb = MUTED

def keybox(doc, text, danger=False):
    bg = rgb_hex(REDL) if danger else rgb_hex(TINT); fg = RED if danger else DARK
    tbl = doc.add_table(rows=1, cols=1); tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]; shade_cell(cell, bg)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text); run.font.name = 'Calibri'; run.font.size = Pt(9)
    run.font.bold = True; run.font.color.rgb = fg
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def dtable(doc, headers, rows, widths=None):
    n = len(headers)
    tbl = doc.add_table(rows=1 + len(rows), cols=n); tbl.style = 'Table Grid'
    for i, h in enumerate(headers):
        shade_cell(tbl.rows[0].cells[i], rgb_hex(MED))
        p = tbl.rows[0].cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(h))
        run.font.name='Calibri'; run.font.size=Pt(8.5); run.font.bold=True; run.font.color.rgb=WHITE
    for ri, row in enumerate(rows):
        bg = rgb_hex(TINT2) if ri % 2 else 'FFFFFF'
        for ci, val in enumerate(row):
            shade_cell(tbl.rows[ri+1].cells[ci], bg)
            p = tbl.rows[ri+1].cells[ci].paragraphs[0]
            txt, bld = (val[0], val[1]) if isinstance(val, tuple) else (val, False)
            run = p.add_run(str(txt)); run.font.bold = bld
            run.font.name='Calibri'; run.font.size=Pt(8.5); run.font.color.rgb=BLACK
    if widths:
        for row in tbl.rows:
            for ci, cell in enumerate(row.cells):
                if ci < len(widths): cell.width = Cm(widths[ci])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ══════════════════════════════════════════════════════════════════════════════
def build_docx():
    doc = Document()
    s = doc.sections[0]
    s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=Cm(2.0); s.top_margin=s.bottom_margin=Cm(1.8)
    doc.core_properties.title  = 'Non-Infectious Pathology in Asian Elephants — Prof. Dr. A B Shrivastava'
    doc.core_properties.author = 'Prof. Dr. A B Shrivastava, NDVSU Jabalpur'

    # TITLE BANNER
    tbl = doc.add_table(rows=1, cols=1); tbl.style='Table Grid'
    cell = tbl.rows[0].cells[0]; shade_cell(cell, rgb_hex(DARK))
    p = cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(20); p.paragraph_format.space_after=Pt(20)
    r0 = p.add_run('NON-INFECTIOUS DISEASE & TOXICOLOGY\n')
    r0.font.name='Calibri'; r0.font.size=Pt(9); r0.font.bold=True; r0.font.color.rgb=LIGHT
    r1 = p.add_run('Non-Infectious Pathology in Asian Elephants\n')
    r1.font.name='Calibri'; r1.font.size=Pt(18); r1.font.bold=True; r1.font.color.rgb=WHITE
    r2 = p.add_run('Comprehensive Study Notes with Slide Images\n')
    r2.font.name='Calibri'; r2.font.size=Pt(10); r2.font.color.rgb=TINT
    r3 = p.add_run('Prof. Dr. A B Shrivastava  ·  Founder Director, School of Wildlife Forensic and Health\n')
    r3.font.name='Calibri'; r3.font.size=Pt(10); r3.font.bold=True; r3.font.color.rgb=WHITE
    r4 = p.add_run('Nanaji Deshmukh Veterinary Science University, Jabalpur 482001, M.P.')
    r4.font.name='Calibri'; r4.font.size=Pt(9); r4.font.italic=True; r4.font.color.rgb=TINT
    doc.add_paragraph()

    keybox(doc,
        'Training Program: Essentials for Mortality Investigation of Asian Elephant\n'
        'Chhattisgarh Forest Department  ·  5–6 June 2026, Raigarh  ·  Technical Support: WII Dehradun\n'
        'Lab Partners: ICAR-IVRI & NDVSU Jabalpur  ·  Notes elaborated with peer-reviewed literature')

    keybox(doc,
        'CONTENTS OF THIS DOCUMENT:\n'
        '  1. Introduction — Non-Infectious Pathology & Elephant Poisoning in India\n'
        '  2. Poisoning Statistics & Categories of Elephant Toxicity\n'
        '  3. Anthropogenic Poisons (Agrochemicals · Cyanide · Fertilizers & Heavy Metals)\n'
        '  4. Mycotoxins & Biological Toxins — Kodo Millet (CPA) · Cyanobacteria · Toxic Plants\n'
        '  5. HCN (Cyanide) Poisoning from Sorghum — Mechanism, Signs, Antidote, Prevention\n'
        '  6. Electrocution in Elephants\n'
        '  7. Pyometra\n'
        '  8. Hyperthermia & Heat Stress\n'
        '  9. Iron Deficiency Anaemia\n'
        '  10. References & Recommended Reading')
    img(doc, 1, 'Slide 1 — Title slide: Non-Infectious Pathology in Asian Elephants; Elephant Poisoning in India', 13)
    page_break(doc)

    # ─── SECTION 1: INTRODUCTION ──────────────────────────────────────────────
    sec1(doc, '1.  Introduction — Non-Infectious Pathology & Elephant Poisoning')

    sec2(doc, '1.1  Non-Infectious Disease in Elephants — Scope')
    para(doc,
        'Non-infectious pathology refers to disease and death NOT caused by transmissible pathogens '
        '(bacteria, viruses, fungi, parasites). In Asian elephants, non-infectious causes are '
        'overwhelmingly anthropogenic (human-caused) and represent the leading cause of unnatural '
        'elephant mortality in India. The major categories are poisoning (intentional and accidental), '
        'electrocution, metabolic/reproductive disorders (pyometra, hyperthermia), and nutritional '
        'deficiency (anaemia).')

    sec2(doc, '1.2  Elephant Poisoning in India — Overview')
    para(doc,
        'Poisoning is a severe threat to pachyderms. In India and across Africa, these incidents '
        'primarily stem from HUMAN-WILDLIFE CONFLICT and ILLEGAL POACHING. A "pachyderm" is a large, '
        'thick-skinned mammal — the term is most commonly used for elephants, rhinoceroses and '
        'hippopotamuses.')
    bullet(doc, 'Poisoning occurs INTENTIONALLY — via illegal poaching (for ivory/body parts) and human-elephant conflict retaliation (crop-raiding revenge)')
    bullet(doc, 'Poisoning also occurs ACCIDENTALLY — ingestion of agrochemicals, contaminated crops, toxic plants, or contaminated water')
    bullet(doc, 'In India, unnatural causes such as electrocution and poisoning claim HUNDREDS of elephant lives every year')
    bullet(doc, 'Elephants are BULK FORAGERS: their high daily food intake (150–300 kg) and water intake (100–200 L) make them highly susceptible to environmental toxins even at low concentrations')
    keybox(doc,
        'KEY CONCEPT: Non-infectious mortality in Indian elephants is largely PREVENTABLE. Unlike '
        'infectious disease, these deaths result from human activity — agrochemical misuse, illegal '
        'electric fencing, retaliatory poisoning, and habitat degradation. Mortality investigation '
        'that correctly identifies the cause is essential for legal action and prevention.')
    page_break(doc)

    # ─── SECTION 2: POISONING STATISTICS ──────────────────────────────────────
    sec1(doc, '2.  Poisoning Statistics & Categories of Toxicity')
    img(doc, 2, 'Slide 2 — Elephant poisoning deaths by state (1997–2005); three categories of toxicity', 13)

    sec2(doc, '2.1  Elephant Deaths by Poisoning Since 1997 (State-Wise)')
    para(doc, 'The following data documents elephant deaths attributed to poisoning across Indian states (Source: WTI Elephant Mortality Database):')
    dtable(doc,
        ['State', 'Years', 'Cases'],
        [('Assam','1997–2005','29'),
         ('Orissa','1997–2003','03'),
         ('West Bengal','1997–2004','03'),
         ('Arunachal Pradesh','1997–2005','02'),
         ('Jharkhand','2002–2003','02'),
         ('Karnataka','1997–2005','02'),
         ('Kerala','1997–2003','02'),
         ('Uttar Pradesh','1997–2004','02'),
         ('Uttaranchal','2001–2004','02'),
         ('Meghalaya','1997–2005','01'),
         ('Tamil Nadu','1997–2005','01'),
         ('TOTAL','1997–2005','49')],
        widths=[6, 5, 3])
    bullet(doc, 'Assam recorded the highest number of poisoning deaths (29 of 49 = 59%) — reflecting intense human-elephant conflict in tea-estate and agricultural landscapes')
    bullet(doc, 'Note: This dataset (1997–2005) represents DOCUMENTED cases only; actual figures are considered substantially higher due to under-reporting and undiagnosed deaths')

    sec2(doc, '2.2  Three Categories of Elephant Toxicity')
    para(doc, 'Elephant poisoning is often due to anthropogenic (human-caused) events and natural toxicities. Elephant toxicity primarily falls into three categories:')
    dtable(doc,
        ['Category', 'Description', 'Examples'],
        [('I. Agricultural / Plant Mycotoxins','Fungal toxins contaminating crops that elephants raid','Cyclopiazonic acid (CPA) on Kodo millet; aflatoxins'),
         ('II. Human-Wildlife Conflict Chemicals','Deliberately or accidentally deployed chemicals','Pesticides (organophosphates, carbamates), cyanide, urea'),
         ('III. Environmental Toxins','Naturally occurring toxins in environment','Cyanobacterial blooms (blue-green algae); toxic plants; heavy metals')],
        widths=[4.5, 6, 6.5])
    keybox(doc,
        'WHY ELEPHANTS ARE HIGHLY SUSCEPTIBLE: Elephants are BULK FORAGERS — they consume enormous '
        'quantities of vegetation (150–300 kg/day) and water (100–200 L/day). This high intake means '
        'that even toxins present at low environmental concentrations can be ingested in lethal total '
        'doses. Their hindgut fermentation also generates and releases toxins (e.g., HCN from cyanogenic '
        'glycosides) over an extended period.')
    page_break(doc)

    # ─── SECTION 3: ANTHROPOGENIC POISONS ─────────────────────────────────────
    sec1(doc, '3.  Anthropogenic Poisons')
    para(doc,
        'Human-elephant conflict may result in intentional or accidental poisonings. The principal '
        'anthropogenic poisons affecting elephants are agrochemicals, cyanide, and fertilizers/heavy metals.')

    sec2(doc, '3.1  Agrochemicals — Pesticides')
    bullet(doc, 'Elephants raiding crops/plantations consume dangerous levels of agrochemicals')
    bullet(doc, 'Most are deadly NEUROTOXINS: organophosphates (e.g., chlorpyrifos, monocrotophos) and carbamates')
    bullet(doc, 'Neonicotinoids (e.g., imidacloprid) — newer systemic insecticides also implicated')
    sec3(doc, 'Organophosphate (OP) / Carbamate Mechanism (web-elaborated):')
    bullet(doc, 'Inhibit acetylcholinesterase (AChE) enzyme → accumulation of acetylcholine at synapses', level=1)
    bullet(doc, 'Cholinergic crisis: SLUDGE signs — Salivation, Lacrimation, Urination, Defecation, GI distress, Emesis', level=1)
    bullet(doc, 'Muscarinic: miosis (pinpoint pupils), bronchorrhoea, bradycardia; Nicotinic: muscle fasciculation, paralysis', level=1)
    bullet(doc, 'Death from respiratory failure (bronchoconstriction + respiratory muscle paralysis + central depression)', level=1)
    bullet(doc, 'ANTIDOTE: Atropine sulphate (blocks muscarinic effects) + Pralidoxime/2-PAM (reactivates AChE — OP only, not carbamate)', level=1)

    sec2(doc, '3.2  Cyanide')
    bullet(doc, 'Used in illegal poaching and mining operations — can cause rapid, fatal hypoxia (cellular asphyxiation)')
    bullet(doc, 'Poachers create "natural salt licks" laced with cyanide near waterholes — elephants attracted to salt are poisoned')
    bullet(doc, 'Mechanism: cyanide binds cytochrome c oxidase → blocks mitochondrial electron transport → cells cannot use oxygen → death')
    bullet(doc, 'See Section 5 for detailed HCN/cyanide poisoning from cyanogenic plants (sorghum)')

    sec2(doc, '3.3  Fertilizers / Heavy Metals')
    bullet(doc, 'Accidental UREA fertilizer toxicity: urea → ammonia in rumen/hindgut → ammonia toxicity → neurological signs, death')
    bullet(doc, 'Heavy metal contamination (ARSENIC, LEAD) can destroy gastrointestinal and renal integrity')
    sec3(doc, 'Heavy Metal Effects (web-elaborated):')
    bullet(doc, 'LEAD: damages nervous system (encephalopathy), GI tract (colic), kidneys; sources — paint, batteries, mining waste', level=1)
    bullet(doc, 'ARSENIC: damages GI mucosa (haemorrhagic enteritis), liver, kidneys; sources — pesticides, industrial effluent', level=1)
    bullet(doc, 'Chronic exposure: weight loss, anaemia, organ failure; Acute: severe GI haemorrhage, rapid death', level=1)
    keybox(doc,
        'TOXICOLOGY SAMPLING (Anthropogenic Poisons): Collect stomach/gut contents, liver, kidney, '
        'brain, and blood — refrigerate or freeze. For OP/carbamate: blood AChE level + gut contents. '
        'For heavy metals: liver and kidney (frozen) for atomic absorption spectroscopy. Always collect '
        'suspected bait/source material. Send to forensic laboratory (IVRI Izatnagar / State FSL) with '
        'a clear suspected-diagnosis note.')
    page_break(doc)

    # ─── SECTION 4: MYCOTOXINS ────────────────────────────────────────────────
    sec1(doc, '4.  Mycotoxins & Biological Toxins')
    img(doc, 3, 'Slide 3 — Mycotoxins (CPA / Kodo millet); cyanotoxins; toxic plants; Bandhavgarh Kodo case', 13)
    img(doc, 4, 'Slide 4 — Field photographs: Kodo millet crop raided; necropsy organ findings', 13)

    sec2(doc, '4.1  Crop-Borne Mycotoxins — Kodo Millet (Cyclopiazonic Acid)')
    para(doc,
        'Elephants raiding farm fields are highly susceptible to mycotoxins like Cyclopiazonic Acid '
        '(CPA). This fungal toxin often infects crops like Kodo millet (Paspalum scrobiculatum) and '
        'can cause acute vascular damage and liver/kidney necrosis.')
    bullet(doc, 'CPA is produced by fungi — primarily Aspergillus and Penicillium species')
    bullet(doc, 'Contamination occurs when heavy or unseasonal rainfall coincides with grain maturation/harvest, creating moist conditions favourable for fungal growth')
    bullet(doc, '"Kodo poisoning" / "Kodua poisoning" is a well-recognised syndrome in livestock and elephants in central India')

    sec2(doc, '4.2  CASE STUDY — Kodo Millet Poisoning, Bandhavgarh Tiger Reserve')
    keybox(doc,
        'RECENT CASE (Madhya Pradesh): A herd of 13 elephants raided the Kodo millet crops near '
        'Bandhavgarh Tiger Reserve. TEN ELEPHANTS DIED. Forensic postmortem and laboratory '
        'investigations revealed that the deaths were caused by accidental toxicity from eating crops '
        'infected with fungal mycotoxins (cyclopiazonic acid). Dr. Shrivastava notes: "I have seen / '
        'treated two cases in KTR (Kanha Tiger Reserve) during 1997–98."', danger=True)
    sec3(doc, 'Management of Kodo Millet Toxicity:')
    bullet(doc, 'Combination of preventive measures in cultivation, storage and processing')
    bullet(doc, 'Biocontrol agents — non-toxigenic fungal strains can competitively exclude toxigenic strains and reduce mycotoxin production')
    bullet(doc, 'Good agricultural practices and proper post-harvest handling: proper drying and airtight storage of Kodo grain')

    sec2(doc, '4.3  Cyanotoxins / Cyanobacteria (Blue-Green Algae)')
    bullet(doc, 'Toxic blue-green algae blooms in stagnant, seasonal water sources produce potent NEUROTOXINS and HEPATOTOXINS')
    bullet(doc, 'These have caused fatal die-offs; mass deaths occasionally linked to toxic algal blooms in watering holes')
    bullet(doc, 'Particularly common during HOT, DRY seasons when water sources shrink and concentrate nutrients')
    sec3(doc, 'Cyanotoxin Types (web-elaborated):')
    bullet(doc, 'MICROCYSTINS (hepatotoxins): produced by Microcystis — cause massive liver haemorrhage and necrosis', level=1)
    bullet(doc, 'ANATOXIN-a / SAXITOXINS (neurotoxins): cause rapid respiratory paralysis and death', level=1)
    bullet(doc, 'Note: The 2020 mass die-off of African elephants in Botswana (Okavango) was attributed to cyanobacterial neurotoxins in waterholes — a landmark example', level=1)

    sec2(doc, '4.4  Toxic Plants')
    bullet(doc, 'Elephants are generally ADEPT at avoiding toxic vegetation through learned behaviour and selective feeding')
    bullet(doc, 'However, ingestion of plants like OLEANDER (Nerium/Thevetia) or certain weeds can cause lethal cardiac or neurological effects')
    bullet(doc, 'Oleander contains cardiac glycosides (oleandrin) → cardiac arrhythmia, heart block, death — even small quantities lethal')
    bullet(doc, 'Some local toxic plants (those containing CALCIUM OXALATE crystals) cause localised severe irritation, swelling, and systemic distress if consumed in large quantities')
    page_break(doc)

    # ─── SECTION 5: HCN POISONING ─────────────────────────────────────────────
    sec1(doc, '5.  HCN (Cyanide) Poisoning from Sorghum')
    img(doc, 5, 'Slide 5 — HCN poisoning: mechanism, signs, risk factors, management & antidote', 13)
    img(doc, 6, 'Slide 6 — HCN poisoning case report (Zoo Print 2010, Dr. Shrivastava); Electrocution intro', 13)

    sec2(doc, '5.1  Overview & Mechanism')
    para(doc,
        'HCN (Hydrocyanic acid or prussic acid) poisoning in elephants occurs when they consume '
        'sorghum that is stressed (drought, frost, or trampling). Ingestion leads to RAPID CELLULAR '
        'ASPHYXIATION.')
    sec3(doc, 'The Mechanism:')
    bullet(doc, 'Sorghum species (Sorghum bicolor) contain a cyanogenic glucoside called DHURRIN')
    bullet(doc, 'When the plant is chewed or its cells are ruptured, plant enzymes (β-glucosidase) mix with dhurrin, releasing HYDROGEN CYANIDE')
    bullet(doc, 'The toxin is rapidly absorbed into the bloodstream and blocks cellular oxygen use (binds cytochrome c oxidase)')
    bullet(doc, 'Result: cells cannot use oxygen despite adequate blood oxygen → death by internal asphyxiation (histotoxic hypoxia)')

    sec2(doc, '5.2  Signs and Symptoms')
    bullet(doc, 'RAPID ONSET: Symptoms appear within minutes to a few hours of consumption')
    bullet(doc, 'BEHAVIOURAL CHANGES: Laboured and rapid breathing, staggering gait, muscle tremors, convulsions, recumbency, and sudden death')
    bullet(doc, 'PHYSICAL SIGNS: A distinct BITTER ALMOND smell; excessive salivation; and BRIGHT CHERRY-RED blood and mucous membranes')
    bullet(doc, 'IMPORTANT DISTINCTION: Cherry-red blood is characteristic of cyanide (oxygen remains bound to haemoglobin); contrast with NITRATE poisoning which causes BROWN (chocolate-coloured) blood from methaemoglobin')

    sec2(doc, '5.3  Risk Factors for Elephants')
    dtable(doc,
        ['Risk Factor', 'Detail'],
        [('Plant Height / Maturity','Young, dark green sorghum regrowth (usually <50 cm tall) holds the HIGHEST concentration of toxins'),
         ('Environmental Stress','Drought-stressed, wilted, or frost-damaged sorghum plants exhibit increased toxicity'),
         ('Digestive Process','Elephants are hindgut fermenters — large quantities of fresh, young, or stressed sorghum fodder quickly overwhelm their metabolic detoxification capacity')],
        widths=[4.5, 12.5])

    sec2(doc, '5.4  Management, Antidote & First Aid')
    sec3(doc, 'Immediate Management:')
    bullet(doc, 'REMOVE THE FOOD SOURCE: move the herd away from the sorghum field immediately')
    bullet(doc, 'ANTIDOTE THERAPY: administer specific cyanide antidote kit (Sodium Nitrite + Sodium Thiosulphate)')
    sec3(doc, 'Antidote & Treatment Protocol:')
    bullet(doc, 'SODIUM NITRITE: administered IV (e.g., 1 gm to 3 gm in 50 ml distilled water) — converts haemoglobin into methaemoglobin, which binds free cyanide to form cyanmethaemoglobin (drawing cyanide off cytochrome oxidase)')
    bullet(doc, 'SODIUM THIOSULPHATE: followed immediately by IV injection — converts cyanide into NON-TOXIC THIOCYANATE, which is safely excreted in urine')
    sec3(doc, 'Prevention:')
    bullet(doc, 'AVOID GRAZING/FEEDING YOUNG PLANTS: do not graze on sorghum less than 18–24 inches (45–60 cm) tall (or <7 weeks old) — concentrations are highest then')
    bullet(doc, 'WAIT AFTER WEATHER STRESS: avoid grazing immediately after a drought ends or a frost occurs — stress spikes the dhurrin content')
    bullet(doc, 'HAY vs PASTURE: Sorghum hay can be safer than fresh pasture (curing reduces HCN by roughly half) — but testing forage before feeding is highly advised')

    sec2(doc, '5.5  CASE STUDY — HCN Poisoning in Circus Elephants (Zoo Print, June 2010)')
    para(doc,
        'Documented by Dr. A B Shrivastava in "Management of Hydrocyanic Acid Poisoning in Asian '
        'Elephants" (Zoo Print vol. XXV, 6 June 2010):')
    bullet(doc, 'Five adult captive elephants were fed a large quantity of fresh Jowar (Sorghum) during their stay at a circus in Jabalpur — these elephants were not accustomed to Sorghum')
    bullet(doc, 'Signs: restlessness and diarrhoea; routine anti-diarrhoeal treatment failed; condition became serious')
    bullet(doc, 'One aged elephant became critical; immediate cessation of Jowar feeding; body temperature 97–98°F (subnormal)')
    bullet(doc, 'Treatment: 8–10 litres dextrose saline + 30 ml MVI via ear vein; Sodium thiosulphate 50 gm in water orally through drinking water, repeated after 12 hrs')
    bullet(doc, 'Aged elephant: given 3rd dose; cold dextrose saline 10 L + small pieces of ice + Tab. Diadin per rectum via tube')
    bullet(doc, 'Outcome: by the next morning the 5th (aged) elephant showed recovery — started drinking water and eating fresh green grass')
    bullet(doc, 'Reference: Radostits OM, Blood DC & Gay CC (1994). Veterinary Medicine, V edition. ELBS/Tindall')
    keybox(doc,
        'HCN TOXICOLOGY SAMPLING: Stomach contents (sealed, airtight, frozen — HCN is volatile and '
        'escapes); blood (10 mL fluoride-oxalate, frozen); liver and muscle (frozen); suspected plant '
        'material. Submit RAPIDLY to toxicology lab as cyanide degrades on storage. Note cherry-red '
        'tissue colour and bitter-almond odour at necropsy as supportive evidence.', danger=True)
    page_break(doc)

    # ─── SECTION 6: ELECTROCUTION ─────────────────────────────────────────────
    sec1(doc, '6.  Electrocution in Elephants')
    img(doc, 7, 'Slide 7 — Electrocution in an elephant (Nov 2020): field photographs and tusk removal', 13)
    img(doc, 9, 'Slide 9 — Electrocution case: carcass lifting, postmortem field photographs', 13)

    sec2(doc, '6.1  Definition & Mechanics')
    para(doc,
        'Electrocution is death or severe injury caused by an electric shock. It occurs when any '
        'individual becomes part of an electrical circuit. High levels of voltage and current disrupt '
        'the body\'s natural electrical systems, commonly triggering cardiac arrest or severe '
        'electrical burns.')
    dtable(doc,
        ['Term', 'Definition'],
        [('Electric Shock','A physical sensation or injury caused by electricity passing through the body (e.g., muscle spasms, minor burns). NON-FATAL injury'),
         ('Electrocution','An electric shock that is FATAL')],
        widths=[4, 13])

    sec2(doc, '6.2  Electrocution in Indian Elephants — Context')
    bullet(doc, 'Electrocution is the LEADING cause of unnatural elephant death in India (~70–75% of unnatural deaths in recent national data)')
    bullet(doc, 'Sources: illegal live wires hung low around crop fields (deliberate, for crop protection or poaching); sagging/low-hanging 11 kV power transmission lines; solar fence tampering')
    bullet(doc, 'Both INTENTIONAL (poaching, conflict retaliation) and ACCIDENTAL (sagging legal power lines) electrocutions occur')

    sec2(doc, '6.3  Postmortem Findings in Electrocution (web-elaborated)')
    bullet(doc, 'ELECTRICAL BURN MARKS (Joule burns): focal, often at the trunk tip, ear, shoulder, or feet — point of contact with the wire; singed skin, charring, crater-like lesions')
    bullet(doc, 'CURRENT MARKS: characteristic blistering/charring at entry and exit points')
    bullet(doc, 'Subcutaneous and intramuscular HAEMORRHAGE along the current path')
    bullet(doc, 'Internal: pulmonary congestion and oedema; petechial haemorrhages on heart (epicardium/endocardium); generalised venous congestion')
    bullet(doc, 'Carcass often found near a known electrical source (fence line, power pole, field boundary)')
    bullet(doc, 'Dr. Shrivastava note: "Experience — charcoal powder" — charred tissue at the contact site has a charcoal-like appearance; histology confirms thermal/electrical coagulative necrosis')
    keybox(doc,
        'ELECTROCUTION INVESTIGATION CHECKLIST:\n'
        '• Document the SCENE: photograph wires, poles, distance from carcass, sagging height\n'
        '• Locate and photograph ENTRY/EXIT burn marks with scale bar\n'
        '• Collect skin from burn margin (histology: coagulative necrosis, "streaming" of epidermal nuclei)\n'
        '• Sample heart, lung (congestion/haemorrhage documentation)\n'
        '• Coordinate with electricity board / police — electrocution is a punishable offence under WPA 1972 & Electricity Act\n'
        '• Distinguish lightning strike (singeing in arborescent/fern-like pattern, multiple animals) from man-made electrocution', danger=True)
    page_break(doc)

    # ─── SECTION 7: PYOMETRA ──────────────────────────────────────────────────
    sec1(doc, '7.  Pyometra')
    img(doc, 8, 'Slide 8 — Pyometra in Asian elephants (clinical photograph; purulent vaginal discharge)', 13)
    img(doc, 10, 'Slide 10 — Vaginal vestibulotomy in Asian elephant: surgical management of pyometra', 13)

    sec2(doc, '7.1  Definition')
    para(doc,
        'Pyometra is a severe, life-threatening uterine infection characterised by the accumulation '
        'of PUS and severe inflammation in the uterus. It is most common in AGING, NULLIPAROUS '
        '(having never given birth) female elephants, and is often linked to prolonged hormonal '
        'fluctuations and underlying reproductive issues.')

    sec2(doc, '7.2  Causes & Pathophysiology')
    bullet(doc, 'HORMONAL CHANGES: Fluctuating progesterone levels and an abnormally thickened uterine lining (endometrium) create an environment favourable for bacterial growth')
    bullet(doc, 'BACTERIAL INFECTION: Bacteria such as Escherichia coli or Enterococcus faecium cause suppurative (pus-forming) inflammation')
    bullet(doc, 'PRE-EXISTING CONDITIONS: Cystic Endometrial Hyperplasia (CEH) or benign tumours (leiomyomas/fibroids) frequently complicate the reproductive tract of older captive elephants and trigger pyometra')
    bullet(doc, 'The non-breeding (nulliparous) state in captive elephants leads to repeated, uninterrupted oestrous cycles — continuous progesterone exposure predisposes to CEH-pyometra complex (web-elaborated)')

    sec2(doc, '7.3  Treatment')
    dtable(doc,
        ['Approach', 'Detail'],
        [('i. Hormonal Downregulation','Use of GnRH vaccines (like Improvac) or Deslorelin implants to suppress reproductive cycles and shrink the affected uterine tissue'),
         ('ii. Aggressive Antibiotic Therapy','Broad-spectrum systemic antibiotics targeting the specific cultured pathogens'),
         ('iii. Uterine Lavage','Flushing the uterus to evacuate accumulated pus'),
         ('iv. Surgical Drainage','Vaginal vestibulotomy — surgical opening to drain the uterine pus (as shown in slide 10)'),
         ('v. Homeopathic (practised)','Pyrogenium 1000 + Hepar sulph 1000 + Secale 1000 (as documented by Dr. Shrivastava)')],
        widths=[4.5, 12.5])
    keybox(doc,
        'CLINICAL NOTE: Pyometra is a reproductive EMERGENCY in elephants. Toxins from the infected '
        'uterus (endotoxaemia) can cause septic shock and death. In captive females, prevention '
        'centres on managed breeding or hormonal cycle suppression to break the cystic-endometrial-'
        'hyperplasia → pyometra cascade. Vaginal vestibulotomy provides surgical drainage when '
        'medical management fails.')
    page_break(doc)

    # ─── SECTION 8: HYPERTHERMIA ──────────────────────────────────────────────
    sec1(doc, '8.  Hyperthermia & Heat Stress')

    sec2(doc, '8.1  Definition')
    para(doc,
        'Hyperthermia is a life-threatening condition caused by high environmental temperature, a '
        'lack of access to water or shade, massive body size, and lack of sweat glands. It is a '
        'serious welfare and mortality concern for both working/captive and free-ranging elephants '
        'during Indian summers.')

    sec2(doc, '8.2  Why Elephants Are Prone to Overheating')
    bullet(doc, 'THE SQUARE-CUBE LAW: As body volume increases, heat production rises faster than the skin surface area available to dissipate it. Massive body size is good for RETAINING heat but POOR at losing it')
    bullet(doc, 'LACK OF SWEAT GLANDS: Elephants do not have traditional (distributed) sweat glands. They have pores only between their toes and rely on "transepidermal water loss" (passive moisture diffusion through skin)')
    bullet(doc, 'METABOLIC HEAT: Continuous physical activity (walking in hot sunlight without rest or water) generates high metabolic heat that easily triggers hyperthermia if the animal cannot cool down')

    sec2(doc, '8.3  Natural Cooling Mechanisms')
    dtable(doc,
        ['Mechanism', 'How It Works'],
        [('The Ears','Elephant ears act as giant THERMAL WINDOWS. Flapping cools the blood circulating through the dense ear vasculature, reducing overall body temperature significantly'),
         ('Wallowing & Spraying','Bathing in mud and water is vital. Water evaporation provides intensive cooling; residual mud protects skin from solar radiation. Elephants also spray their own saliva'),
         ('Skin Permeability','The elephant\'s hide becomes more PERMEABLE in hot weather, allowing them to lose moisture and cool down faster through evaporation — provided they have access to drinking water'),
         ('Skin Wrinkles','(web-elaborated) Deep skin wrinkles increase surface area and retain water/mud after wallowing, prolonging evaporative cooling')],
        widths=[4, 13])

    sec2(doc, '8.4  Prevention & Management (web-elaborated)')
    bullet(doc, 'Ensure constant access to SHADE and clean DRINKING WATER, especially 11 AM – 4 PM')
    bullet(doc, 'Provide wallowing pools / mud baths for captive elephants')
    bullet(doc, 'Avoid working/walking elephants during peak heat; reschedule to early morning or evening')
    bullet(doc, 'Recognise heat stress signs: rapid breathing, ear-flapping intensified, lethargy, reluctance to move, collapse')
    bullet(doc, 'Emergency cooling: continuous water spraying (especially ears, head), move to shade, IV fluids')
    page_break(doc)

    # ─── SECTION 9: IRON DEFICIENCY ANAEMIA ───────────────────────────────────
    sec1(doc, '9.  Iron Deficiency Anaemia')
    img(doc, 11, 'Slide 11 — Iron Deficiency Anaemia: blood smear showing microcytic hypochromic target cells', 13)

    sec2(doc, '9.1  Definition')
    para(doc,
        'Anaemia is a qualitative and quantitative reduction in the quality (and oxygen-carrying '
        'capacity) of the blood. Iron deficiency anaemia results from inadequate iron for haemoglobin '
        'synthesis, producing characteristic microcytic, hypochromic red blood cells.')

    sec2(doc, '9.2  Clinical Signs Observed')
    bullet(doc, 'Pale mucous membranes (conjunctiva, oral mucosa)')
    bullet(doc, 'Extreme weakness; unable to walk; very slow walk and reduced activity')
    bullet(doc, 'Lethargy, exercise intolerance, reduced work capacity in captive elephants')

    sec2(doc, '9.3  Diagnosis')
    bullet(doc, 'BLOOD EXAMINATION: Low Haemoglobin (Hb)')
    bullet(doc, 'BLOOD SMEAR: typical TARGET CELLS (codocytes) — RBCs that are MICROCYTIC (small) and HYPOCHROMIC (pale, with increased central pallor)')
    bullet(doc, 'Target cell appearance is characteristic of iron deficiency (and also seen in liver disease, thalassaemia in other species)')
    bullet(doc, 'Reference normal Hb range (Indian captive elephants, web-elaborated): ~8.6–16.8 g/dL; values below this range confirm anaemia')

    sec2(doc, '9.4  Treatment')
    sec3(doc, '1. Diet Change — to full normal elephant diet:')
    bullet(doc, 'Soybeans, horse gram, Napier grass, bamboo, leaves, grains, fruits', level=1)
    sec3(doc, '2. Supplementation:')
    bullet(doc, 'IRON SUPPLEMENTS: Jaggery (gur), honey — traditional iron-rich supplements', level=1)
    bullet(doc, 'VITAMIN B12 INJECTIONS: supports erythropoiesis (red blood cell production)', level=1)
    bullet(doc, 'MULTIVITAMINS, Mineral mixture (M mixture)', level=1)
    keybox(doc,
        'INVESTIGATION NOTE: Iron deficiency anaemia in elephants is often SECONDARY — investigate '
        'the underlying cause: chronic blood loss (haematophagous parasites — ticks, lice, GI '
        'helminths such as Bathmostomum), nutritional inadequacy, or chronic disease. Treating only '
        'the anaemia without addressing the cause leads to relapse. Always run a parasitological '
        'examination (faecal egg count, ectoparasite check) alongside iron supplementation.')
    page_break(doc)

    # ─── SECTION 10: REFERENCES ───────────────────────────────────────────────
    sec1(doc, '10.  References & Recommended Reading')

    sec2(doc, '10.1  Works by Prof. Dr. A B Shrivastava')
    bullet(doc, 'Shrivastav AB (2010). Management of Hydrocyanic Acid Poisoning in Asian Elephants. Zoo Print vol. XXV, 6 June 2010')
    bullet(doc, 'Shrivastav AB & Sharma RK (2008). A Manual of Wildlife Health Management in Protected Areas. CVAH, JNKVV, Jabalpur')
    bullet(doc, 'Nigam P, Habib J & Pandey H (Ed.) (2021). Caring for Elephants: Managing Health & Welfare in Captivity')
    bullet(doc, 'MoEFCC/WII (2020). Necropsy and Carcass Disposal of Asian Elephant: Recommended Operating Procedure')

    sec2(doc, '10.2  Peer-Reviewed & Reference Literature')
    bullet(doc, 'Radostits OM, Blood DC & Gay CC (1994/2007). Veterinary Medicine: A Textbook of the Diseases of Cattle, Sheep, Pigs, Goats and Horses. ELBS/Tindall')
    bullet(doc, 'Fowler ME & Mikota SK (2006). Biology, Medicine and Surgery of Elephants. Blackwell, Iowa')
    bullet(doc, 'WTI (Wildlife Trust of India). Elephant Mortality Database — poisoning records 1997–2005')
    bullet(doc, 'Bengis RG et al. (2004). The role of wildlife in emerging and re-emerging zoonoses. Rev Sci Tech OIE')
    bullet(doc, 'Wang DZ (2008). Neurotoxins from marine and freshwater algae. Marine Drugs 6(2): 349–371')
    bullet(doc, 'Hooijberg EH et al. (2023). Haematology reference intervals for Asian elephants. (Indian captive elephant blood parameters)')
    bullet(doc, 'Mahato G et al. (2021). Heavy metal toxicosis in wild and captive elephants — review. Indian Journal of Veterinary Pathology')
    bullet(doc, 'Project Elephant, MoEFCC (2017). Gajah: Securing the Future for Elephants in India')

    para(doc,
        'Notes prepared from the presentation of Prof. Dr. A B Shrivastava (06-06-2026), elaborated '
        'with peer-reviewed literature and current guidelines. Training Program: Essentials for '
        'Mortality Investigation of Asian Elephant, Chhattisgarh Forest Department, 5–6 June 2026, Raigarh.',
        size=8, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(OUT_D)
    print(f'  ✓  {OUT_D}  ({os.path.getsize(OUT_D)//1024} KB)')


# ══════════════════════════════════════════════════════════════════════════════
#  BUILD PDF
# ══════════════════════════════════════════════════════════════════════════════
def build_pdf():
    W, H = A4
    doc = SimpleDocTemplate(OUT_P, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=1.8*cm, bottomMargin=1.8*cm)
    styles = getSampleStyleSheet(); story = []
    DK = colors.HexColor('#1B4332'); MD = colors.HexColor('#2D6A4F')
    LT = colors.HexColor('#52B788'); TN = colors.HexColor('#D8F3DC')
    T2 = colors.HexColor('#EAF7EE'); RD = colors.HexColor('#FAE0D8')
    WH = colors.white; MT = colors.HexColor('#555555'); DRD = colors.HexColor('#8B0000')

    def ST(name, **kw):
        s = ParagraphStyle(name, parent=styles['Normal'])
        for k, v in kw.items(): setattr(s, k, v)
        return s
    sH1 = ST('H1', fontSize=12, textColor=WH, leading=16, fontName='Helvetica-Bold')
    sH2 = ST('H2', fontSize=10, textColor=WH, leading=13, fontName='Helvetica-Bold')
    sH3 = ST('H3', fontSize=9.5, textColor=MD, leading=13, fontName='Helvetica-Bold', spaceAfter=2)
    sBD = ST('BD', fontSize=9, leading=13, spaceAfter=3, alignment=TA_JUSTIFY)
    sBL = ST('BL', fontSize=9, leading=13, leftIndent=12, firstLineIndent=-9, spaceAfter=2)
    sBL2= ST('B2', fontSize=8.5, leading=12, leftIndent=24, firstLineIndent=-9, spaceAfter=1)
    sCAP= ST('CP', fontSize=7.5, textColor=MT, leading=10, alignment=TA_CENTER, fontName='Helvetica-Oblique', spaceAfter=4)
    sKY = ST('KY', fontSize=8.5, textColor=DK, leading=12, fontName='Helvetica-Bold')
    sKYR= ST('KR', fontSize=8.5, textColor=DRD, leading=12, fontName='Helvetica-Bold')
    sFT = ST('FT', fontSize=7.5, textColor=MT, alignment=TA_CENTER, fontName='Helvetica-Oblique')
    sTD = ST('TD', fontSize=8, leading=11, spaceAfter=1)
    sTH = ST('TH', fontSize=8, leading=11, textColor=WH, fontName='Helvetica-Bold')
    sTIT= ST('TT', fontSize=17, textColor=WH, leading=23, fontName='Helvetica-Bold', alignment=TA_CENTER)
    sSUB= ST('SB', fontSize=10, textColor=LT, leading=14, fontName='Helvetica-Oblique', alignment=TA_CENTER)
    sSB2= ST('SC', fontSize=9, textColor=WH, leading=12, fontName='Helvetica-Bold', alignment=TA_CENTER)
    CW = W - 4*cm

    def h1(txt):
        story.append(Spacer(1, 0.2*cm))
        t = Table([[Paragraph(txt, sH1)]], colWidths=[CW])
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),DK),('TOPPADDING',(0,0),(-1,-1),7),
                               ('BOTTOMPADDING',(0,0),(-1,-1),7),('LEFTPADDING',(0,0),(-1,-1),8)]))
        story.append(t); story.append(Spacer(1, 0.12*cm))
    def h2(txt):
        story.append(Spacer(1, 0.1*cm))
        t = Table([[Paragraph(txt, sH2)]], colWidths=[CW])
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),MD),('TOPPADDING',(0,0),(-1,-1),3),
                               ('BOTTOMPADDING',(0,0),(-1,-1),3),('LEFTPADDING',(0,0),(-1,-1),6)]))
        story.append(t); story.append(Spacer(1, 0.08*cm))
    def h3(txt): story.append(Paragraph(txt, sH3))
    def pb(txt, sub=False):
        prefix = '    ◦  ' if sub else '•  '
        story.append(Paragraph(prefix + txt, sBL2 if sub else sBL))
    def bdy(txt): story.append(Paragraph(txt, sBD))
    def kyb(txt, danger=False):
        ks = sKYR if danger else sKY; bg = RD if danger else TN
        t = Table([[Paragraph(txt, ks)]], colWidths=[CW])
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),bg),('TOPPADDING',(0,0),(-1,-1),5),
                               ('BOTTOMPADDING',(0,0),(-1,-1),5),('LEFTPADDING',(0,0),(-1,-1),8),
                               ('BOX',(0,0),(-1,-1),0.5,MD)]))
        story.append(t); story.append(Spacer(1, 0.1*cm))
    def pi(n, cap='', w=13):
        path = slide(n)
        if os.path.exists(path):
            story.append(Spacer(1, 0.15*cm))
            ri = RLImage(path, width=w*cm, height=w*cm*0.78); ri.hAlign = 'CENTER'
            story.append(ri)
            if cap: story.append(Paragraph(cap, sCAP))
            story.append(Spacer(1, 0.1*cm))
    def dtbl(hdrs, rows, wds=None):
        n = len(hdrs)
        if not wds: wds = [17/n]*n
        data = [[Paragraph(str(h), sTH) for h in hdrs]]
        for row in rows:
            data.append([Paragraph(str(v[0] if isinstance(v, tuple) else v), sTD) for v in row])
        t = Table(data, colWidths=[w*cm for w in wds])
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),MD),('ROWBACKGROUNDS',(0,1),(-1,-1),[WH,T2]),
            ('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#BBBBBB')),('TOPPADDING',(0,0),(-1,-1),3),
            ('BOTTOMPADDING',(0,0),(-1,-1),3),('LEFTPADDING',(0,0),(-1,-1),4),('VALIGN',(0,0),(-1,-1),'TOP')]))
        story.append(t); story.append(Spacer(1, 0.15*cm))
    def hr(): story.append(HRFlowable(width='100%', thickness=0.5, color=LT, spaceAfter=4))

    # TITLE
    tb = Table([[Paragraph('NON-INFECTIOUS DISEASE & TOXICOLOGY', sSB2)],
                [Paragraph('Non-Infectious Pathology in Asian Elephants', sTIT)],
                [Paragraph('Comprehensive Study Notes with Slide Images', sSUB)],
                [Paragraph('Prof. Dr. A B Shrivastava  ·  School of Wildlife Forensic and Health', sSUB)],
                [Paragraph('Nanaji Deshmukh Veterinary Science University, Jabalpur 482001, M.P.', sSUB)]],
               colWidths=[CW])
    tb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),DK),('TOPPADDING',(0,0),(-1,-1),12),
                            ('BOTTOMPADDING',(0,0),(-1,-1),12),('LEFTPADDING',(0,0),(-1,-1),12)]))
    story.append(Spacer(1, 0.3*cm)); story.append(tb); story.append(Spacer(1, 0.25*cm))
    kyb('Training Program: Essentials for Mortality Investigation of Asian Elephant\n'
        'Chhattisgarh Forest Department  ·  5–6 June 2026, Raigarh  ·  Technical Support: WII Dehradun')
    kyb('CONTENTS: 1 Intro  |  2 Poisoning Stats  |  3 Anthropogenic Poisons  |  4 Mycotoxins  |  '
        '5 HCN/Cyanide  |  6 Electrocution  |  7 Pyometra  |  8 Hyperthermia  |  9 Anaemia  |  10 References')
    pi(1, 'Slide 1 — Title slide; Elephant Poisoning in India introduction', 12)
    story.append(PageBreak())

    # S1
    h1('1.  Introduction — Non-Infectious Pathology & Elephant Poisoning')
    h2('1.1  Scope')
    bdy('Non-infectious pathology = disease/death NOT caused by transmissible pathogens. In Asian elephants these causes are overwhelmingly ANTHROPOGENIC and are the leading cause of unnatural elephant mortality in India: poisoning, electrocution, metabolic/reproductive disorders, and nutritional deficiency.')
    h2('1.2  Elephant Poisoning in India')
    bdy('Poisoning is a severe threat to pachyderms. Incidents primarily stem from HUMAN-WILDLIFE CONFLICT and ILLEGAL POACHING — intentionally (poaching, retaliation) or accidentally (agrochemicals, contaminated crops/water).')
    pb('Elephants are BULK FORAGERS — 150–300 kg food + 100–200 L water/day → highly susceptible to environmental toxins')
    pb('Unnatural causes (electrocution, poisoning) claim hundreds of elephant lives yearly in India')
    kyb('KEY: Non-infectious elephant mortality is largely PREVENTABLE — it results from human activity. Correct cause-of-death identification is essential for legal action and prevention.')
    story.append(PageBreak())

    # S2
    h1('2.  Poisoning Statistics & Categories of Toxicity')
    pi(2, 'Slide 2 — State-wise poisoning deaths (1997–2005, total 49); three categories of toxicity', 12)
    h2('2.1  Elephant Deaths by Poisoning, State-Wise (WTI Database)')
    dtbl(['State','Years','Cases'],
         [('Assam','1997–2005','29'),('Orissa','1997–2003','03'),('West Bengal','1997–2004','03'),
          ('Arunachal Pradesh','1997–2005','02'),('Jharkhand','2002–2003','02'),('Karnataka','1997–2005','02'),
          ('Kerala','1997–2003','02'),('Uttar Pradesh','1997–2004','02'),('Uttaranchal','2001–2004','02'),
          ('Meghalaya','1997–2005','01'),('Tamil Nadu','1997–2005','01'),('TOTAL','1997–2005','49')],
         wds=[7,6,4])
    pb('Assam = 29/49 (59%) — intense human-elephant conflict in tea-estate/agricultural landscapes')
    pb('Documented cases only — actual figures considered substantially higher (under-reporting)')
    h2('2.2  Three Categories of Elephant Toxicity')
    dtbl(['Category','Description','Examples'],
         [('I. Agricultural / Plant Mycotoxins','Fungal toxins on raided crops','Cyclopiazonic acid (CPA) on Kodo millet; aflatoxins'),
          ('II. Human-Wildlife Conflict Chemicals','Deliberate/accidental chemicals','Pesticides (organophosphates), cyanide, urea'),
          ('III. Environmental Toxins','Naturally occurring toxins','Cyanobacteria (blue-green algae); toxic plants; heavy metals')],
         wds=[4.5,6,6.5])
    kyb('SUSCEPTIBILITY: Bulk foraging (150–300 kg/day) means even low-concentration toxins reach lethal total doses. Hindgut fermentation also releases toxins (e.g. HCN) over an extended period.')
    story.append(PageBreak())

    # S3
    h1('3.  Anthropogenic Poisons')
    h2('3.1  Agrochemicals — Pesticides (Neurotoxins)')
    pb('Crop-raiding elephants consume organophosphates (chlorpyrifos, monocrotophos), carbamates, neonicotinoids (imidacloprid)')
    pb('OP/Carbamate mechanism: inhibit acetylcholinesterase → acetylcholine accumulates → cholinergic crisis')
    pb('SLUDGE signs: Salivation, Lacrimation, Urination, Defecation, GI distress, Emesis; miosis; muscle fasciculation')
    pb('Death from respiratory failure; ANTIDOTE: Atropine + Pralidoxime (2-PAM, for OP only)')
    h2('3.2  Cyanide')
    pb('Used in illegal poaching/mining → rapid fatal hypoxia; poachers lace salt licks near waterholes')
    pb('Binds cytochrome c oxidase → blocks mitochondrial oxygen use (see Section 5 for detail)')
    h2('3.3  Fertilizers / Heavy Metals')
    pb('UREA toxicity → ammonia → neurological signs, death')
    pb('LEAD: encephalopathy, GI colic, kidney damage (paint, batteries, mining waste)')
    pb('ARSENIC: haemorrhagic enteritis, liver/kidney damage (pesticides, industrial effluent)')
    kyb('SAMPLING: gut contents, liver, kidney, brain, blood (frozen). OP/carbamate → blood AChE. Heavy metals → liver/kidney for AAS. Always collect suspected bait. Send to IVRI/State FSL with suspected-diagnosis note.')
    story.append(PageBreak())

    # S4
    h1('4.  Mycotoxins & Biological Toxins')
    pi(3, 'Slide 3 — Mycotoxins (CPA/Kodo); cyanotoxins; toxic plants; Bandhavgarh Kodo case', 12)
    pi(4, 'Slide 4 — Field photographs: raided Kodo millet crop; necropsy organ findings', 12)
    h2('4.1  Crop-Borne Mycotoxins — Kodo Millet (Cyclopiazonic Acid)')
    bdy('Elephants raiding farm fields are highly susceptible to mycotoxins like Cyclopiazonic Acid (CPA) — a fungal toxin (Aspergillus, Penicillium) that infects Kodo millet (Paspalum scrobiculatum) and causes acute vascular damage and liver/kidney necrosis. Contamination occurs when heavy/unseasonal rainfall coincides with grain maturation.')
    kyb('CASE — Bandhavgarh Tiger Reserve (MP): A herd of 13 elephants raided Kodo millet crops; TEN ELEPHANTS DIED. Forensic PM + lab confirmed accidental mycotoxin (CPA) toxicity. Dr. Shrivastava: "I have seen/treated two cases in KTR during 1997–98."', danger=True)
    pb('Management: preventive cultivation/storage/processing; biocontrol (non-toxigenic fungal strains); proper drying + airtight storage of Kodo')
    h2('4.2  Cyanotoxins / Cyanobacteria (Blue-Green Algae)')
    pb('Toxic blue-green algae blooms in stagnant seasonal water → neurotoxins + hepatotoxins → fatal die-offs')
    pb('Common in hot, dry seasons when water sources shrink/concentrate')
    pb('Microcystins (hepatotoxins, Microcystis) → liver necrosis; Anatoxin-a/Saxitoxins (neurotoxins) → respiratory paralysis')
    pb('2020 Botswana (Okavango) African elephant mass die-off attributed to cyanobacterial neurotoxins')
    h2('4.3  Toxic Plants')
    pb('Elephants generally avoid toxic vegetation, but oleander (cardiac glycosides) → fatal arrhythmia')
    pb('Calcium oxalate-containing plants → severe oral/GI irritation, swelling, systemic distress')
    story.append(PageBreak())

    # S5
    h1('5.  HCN (Cyanide) Poisoning from Sorghum')
    pi(5, 'Slide 5 — HCN poisoning: mechanism, signs, risk factors, management & antidote', 12)
    pi(6, 'Slide 6 — HCN case report (Zoo Print 2010); Electrocution introduction', 12)
    h2('5.1  Mechanism')
    bdy('HCN (prussic acid) poisoning occurs when elephants consume STRESSED sorghum (drought, frost, trampling) → rapid cellular asphyxiation. Sorghum (Sorghum bicolor) contains the cyanogenic glucoside DHURRIN. When chewed, plant enzymes mix with dhurrin → release hydrogen cyanide → binds cytochrome c oxidase → cells cannot use oxygen → histotoxic hypoxia → death.')
    h2('5.2  Signs & Symptoms')
    pb('RAPID ONSET: minutes to a few hours after consumption')
    pb('Laboured/rapid breathing, staggering gait, muscle tremors, convulsions, recumbency, sudden death')
    pb('BITTER ALMOND smell; excessive salivation; BRIGHT CHERRY-RED blood and mucous membranes')
    pb('DISTINCTION: cyanide → cherry-red blood; nitrate poisoning → brown (chocolate) blood')
    h2('5.3  Risk Factors')
    dtbl(['Risk Factor','Detail'],
         [('Plant height/maturity','Young dark-green regrowth (<50 cm tall) = highest toxin concentration'),
          ('Environmental stress','Drought-stressed, wilted, or frost-damaged plants are most toxic'),
          ('Digestive process','Hindgut fermenters; large quantities overwhelm detoxification capacity')],
         wds=[4.5,12.5])
    h2('5.4  Antidote & Prevention')
    pb('Remove food source immediately; move herd away from sorghum field')
    pb('SODIUM NITRITE IV (1–3 gm in 50 ml distilled water) → forms methaemoglobin → binds cyanide as cyanmethaemoglobin')
    pb('SODIUM THIOSULPHATE IV → converts cyanide to non-toxic thiocyanate → excreted in urine')
    pb('Prevention: avoid young plants (<45–60 cm); wait after drought/frost; hay safer than fresh pasture (test forage)')
    kyb('CASE — Circus elephants, Jabalpur (Zoo Print 2010, Dr. Shrivastava): 5 adults fed fresh Jowar (unaccustomed) → restlessness, diarrhoea, subnormal temp 97–98°F. Treated: dextrose saline + MVI; sodium thiosulphate 50 gm oral, repeated 12 hrs; cold saline + ice + Diadin per rectum. Aged elephant recovered next morning.', danger=True)
    story.append(PageBreak())

    # S6
    h1('6.  Electrocution in Elephants')
    pi(7, 'Slide 7 — Electrocution in an elephant (Nov 2020): field photographs', 12)
    pi(9, 'Slide 9 — Electrocution case: carcass handling and postmortem photographs', 12)
    h2('6.1  Definition & Mechanics')
    bdy('Electrocution = death/severe injury from electric shock; occurs when the animal becomes part of an electrical circuit. High voltage/current disrupts the body\'s electrical systems → cardiac arrest or severe electrical burns.')
    dtbl(['Term','Definition'],
         [('Electric Shock','Physical sensation/injury from electricity (muscle spasm, minor burns) — NON-FATAL'),
          ('Electrocution','An electric shock that is FATAL')],
         wds=[4,13])
    h2('6.2  Context in India')
    pb('LEADING cause of unnatural elephant death in India (~70–75% of unnatural deaths)')
    pb('Sources: illegal low live wires around crops; sagging 11 kV power lines; solar-fence tampering')
    pb('Both intentional (poaching/retaliation) and accidental (sagging legal lines)')
    h2('6.3  Postmortem Findings')
    pb('Electrical BURN MARKS (Joule burns) at contact point — trunk tip, ear, shoulder, feet; charring, crater lesions')
    pb('Subcutaneous/intramuscular haemorrhage along current path; entry & exit current marks')
    pb('Internal: pulmonary congestion/oedema; cardiac petechiae; generalised venous congestion')
    pb('Carcass near electrical source; Dr. Shrivastava note: "charcoal powder" appearance of charred contact tissue')
    kyb('INVESTIGATION: photograph scene (wires, poles, sag height, distance); burn marks with scale bar; skin histology (coagulative necrosis, epidermal nuclear streaming); coordinate with electricity board/police (punishable under WPA 1972 & Electricity Act); distinguish lightning (fern-pattern singeing, multiple animals).', danger=True)
    story.append(PageBreak())

    # S7
    h1('7.  Pyometra')
    pi(8, 'Slide 8 — Pyometra in Asian elephants (clinical photograph; purulent discharge)', 12)
    pi(10, 'Slide 10 — Vaginal vestibulotomy in Asian elephant: surgical drainage', 12)
    h2('7.1  Definition')
    bdy('Pyometra = severe, life-threatening uterine infection with accumulation of PUS and severe inflammation. Most common in AGING, NULLIPAROUS (never-bred) females, linked to prolonged hormonal fluctuation and reproductive issues.')
    h2('7.2  Causes & Pathophysiology')
    pb('Hormonal: fluctuating progesterone + thickened endometrium favour bacterial growth')
    pb('Bacterial: E. coli or Enterococcus faecium → suppurative inflammation')
    pb('Pre-existing: Cystic Endometrial Hyperplasia (CEH) or leiomyomas in older captive elephants trigger pyometra')
    pb('Nulliparous state → repeated uninterrupted cycles → continuous progesterone → CEH-pyometra complex')
    h2('7.3  Treatment')
    dtbl(['Approach','Detail'],
         [('Hormonal downregulation','GnRH vaccines (Improvac) or Deslorelin implants to suppress cycles, shrink tissue'),
          ('Antibiotic therapy','Broad-spectrum systemic antibiotics targeting cultured pathogens'),
          ('Uterine lavage','Flushing the uterus to evacuate pus'),
          ('Surgical drainage','Vaginal vestibulotomy (see slide 10)'),
          ('Homeopathic (practised)','Pyrogenium 1000 + Hepar sulph 1000 + Secale 1000 (Dr. Shrivastava)')],
         wds=[4.5,12.5])
    kyb('Pyometra is a reproductive EMERGENCY — uterine endotoxaemia can cause septic shock and death. Prevention in captive females: managed breeding or hormonal cycle suppression to break the CEH→pyometra cascade.')
    story.append(PageBreak())

    # S8
    h1('8.  Hyperthermia & Heat Stress')
    h2('8.1  Definition')
    bdy('Hyperthermia is a life-threatening condition from high environmental temperature, lack of access to water/shade, massive body size, and lack of sweat glands — a serious mortality concern during Indian summers.')
    h2('8.2  Why Elephants Overheat')
    pb('SQUARE-CUBE LAW: heat production rises faster than skin surface area; massive size retains heat but loses it poorly')
    pb('NO SWEAT GLANDS: pores only between toes; rely on transepidermal water loss')
    pb('METABOLIC HEAT: continuous walking in hot sun without rest/water triggers hyperthermia')
    h2('8.3  Natural Cooling Mechanisms')
    dtbl(['Mechanism','How It Works'],
         [('The Ears','Giant thermal windows; flapping cools blood in dense ear vasculature'),
          ('Wallowing & Spraying','Mud/water evaporation cools; residual mud blocks solar radiation; saliva spraying'),
          ('Skin Permeability','Hide becomes more permeable in heat → faster evaporative cooling (needs drinking water)'),
          ('Skin Wrinkles','Increase surface area; retain water/mud → prolong evaporative cooling')],
         wds=[4,13])
    h2('8.4  Prevention & Management')
    pb('Constant access to SHADE + clean DRINKING WATER, especially 11 AM–4 PM')
    pb('Provide wallowing pools/mud baths; avoid working elephants during peak heat')
    pb('Recognise signs: rapid breathing, intensified ear-flapping, lethargy, collapse')
    pb('Emergency cooling: continuous water spray (ears/head), shade, IV fluids')
    story.append(PageBreak())

    # S9
    h1('9.  Iron Deficiency Anaemia')
    pi(11, 'Slide 11 — Blood smear: microcytic hypochromic target cells (codocytes)', 12)
    h2('9.1  Definition & Signs')
    bdy('Anaemia = qualitative and quantitative reduction in blood quality / oxygen-carrying capacity. Iron deficiency → microcytic, hypochromic RBCs. Signs: pale mucous membranes, extreme weakness, unable to walk, very slow walk and reduced activity.')
    h2('9.2  Diagnosis')
    pb('Blood examination: LOW Haemoglobin (Hb)')
    pb('Blood smear: typical TARGET CELLS (codocytes) — MICROCYTIC, HYPOCHROMIC RBCs')
    pb('Reference Hb (Indian captive elephants): ~8.6–16.8 g/dL; below range confirms anaemia')
    h2('9.3  Treatment')
    pb('Diet change to full normal elephant diet: soybeans, horse gram, Napier grass, bamboo, leaves, grains, fruits')
    pb('Iron supplements: Jaggery (gur), honey')
    pb('Vitamin B12 injections; Multivitamins + mineral mixture')
    kyb('INVESTIGATION: Iron deficiency anaemia is often SECONDARY — investigate chronic blood loss (ticks, lice, GI helminths), nutritional inadequacy, or chronic disease. Run a parasitological exam alongside iron supplementation to prevent relapse.')
    story.append(PageBreak())

    # S10
    h1('10.  References & Recommended Reading')
    h2('10.1  Works by Prof. Dr. A B Shrivastava')
    pb('Shrivastav AB (2010). Management of Hydrocyanic Acid Poisoning in Asian Elephants. Zoo Print vol. XXV, 6 June 2010')
    pb('Shrivastav AB & Sharma RK (2008). A Manual of Wildlife Health Management in Protected Areas. CVAH, JNKVV, Jabalpur')
    pb('Nigam P, Habib J & Pandey H (Ed.) (2021). Caring for Elephants: Managing Health & Welfare in Captivity')
    pb('MoEFCC/WII (2020). Necropsy and Carcass Disposal of Asian Elephant: Recommended Operating Procedure')
    h2('10.2  Peer-Reviewed & Reference Literature')
    pb('Radostits OM, Blood DC & Gay CC (1994/2007). Veterinary Medicine. ELBS/Tindall')
    pb('Fowler ME & Mikota SK (2006). Biology, Medicine and Surgery of Elephants. Blackwell, Iowa')
    pb('WTI. Elephant Mortality Database — poisoning records 1997–2005')
    pb('Wang DZ (2008). Neurotoxins from marine and freshwater algae. Marine Drugs 6(2): 349–371')
    pb('Mahato G et al. (2021). Heavy metal toxicosis in elephants — review. Indian J Vet Pathology')
    pb('Project Elephant, MoEFCC (2017). Gajah: Securing the Future for Elephants in India')
    hr()
    story.append(Paragraph(
        'Prof. Dr. A B Shrivastava, NDVSU Jabalpur  |  Presentation: 06-06-2026  |  '
        'Training Program: Mortality Investigation of Asian Elephant, Raigarh, Chhattisgarh  |  '
        'Notes elaborated with peer-reviewed literature', sFT))

    doc.build(story)
    print(f'  ✓  {OUT_P}  ({os.path.getsize(OUT_P)//1024} KB)')


if __name__ == '__main__':
    print('Building ABS Non-Infectious Pathology Notes...\n')
    print('  → Word document...'); build_docx()
    print('  → PDF document...'); build_pdf()
    print('\nDone.')
